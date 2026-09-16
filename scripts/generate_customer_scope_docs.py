from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "customer-package"
OUT.mkdir(parents=True, exist_ok=True)

PROJECT = "TIMEHOUSE"
PROJECT_LONG = "Hệ thống quản lý vận hành nhà cho thuê"
VERSION = "1.0"
DATE = "14/09/2026"
CLIENT = "[TÊN KHÁCH HÀNG]"
VENDOR = "[TÊN ĐƠN VỊ THỰC HIỆN]"

NAVY = "16324F"
BLUE = "1769AA"
GREEN = "2B7A66"
PURPLE = "6B4EA0"
ORANGE = "E48732"
LIGHT_BLUE = "EAF3F9"
LIGHT_GREEN = "EAF5F1"
LIGHT_PURPLE = "F1EDF8"
LIGHT_GRAY = "F2F4F6"
MID_GRAY = "D7DDE3"
DARK_GRAY = "3F4A54"
WHITE = "FFFFFF"
RED = "B42318"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, **edges):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "start", "bottom", "end", "insideH", "insideV"):
        if edge not in edges:
            continue
        edge_data = edges[edge]
        tag = "left" if edge == "start" else "right" if edge == "end" else edge
        node = borders.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            borders.append(node)
        for key in ("val", "sz", "space", "color"):
            if key in edge_data:
                node.set(qn(f"w:{key}"), str(edge_data[key]))


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_repeat_header_text(run):
    run.bold = True
    run.font.color.rgb = RGBColor(255, 255, 255)
    run.font.size = Pt(9)


def add_field(run, field_code):
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = field_code
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Trang ")
    add_field(run, "PAGE")
    run.add_text(" / ")
    add_field(run, "NUMPAGES")


def set_update_fields(doc):
    settings = doc.settings.element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def set_repeat_heading(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    keep_next = OxmlElement("w:keepNext")
    p_pr.append(keep_next)


def configure_styles(doc, accent):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string(DARK_GRAY)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.12

    for style_name, size, color in (
        ("Title", 30, NAVY),
        ("Subtitle", 13, DARK_GRAY),
        ("Heading 1", 17, accent),
        ("Heading 2", 12, NAVY),
        ("Heading 3", 10.5, DARK_GRAY),
    ):
        style = styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        if style_name.startswith("Heading"):
            style.font.bold = True
            style.paragraph_format.space_before = Pt(12)
            style.paragraph_format.space_after = Pt(5)
            style.paragraph_format.keep_with_next = True

    if "TH Small" not in styles:
        small = styles.add_style("TH Small", WD_STYLE_TYPE.PARAGRAPH)
    else:
        small = styles["TH Small"]
    small.font.name = "Arial"
    small._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    small.font.size = Pt(8)
    small.font.color.rgb = RGBColor.from_string(DARK_GRAY)
    small.paragraph_format.space_after = Pt(2)


def configure_document(doc, title, accent):
    configure_styles(doc, accent)
    set_update_fields(doc)
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(0.7)
    section.footer_distance = Cm(0.7)

    header = section.header
    p = header.paragraphs[0]
    p.text = f"{PROJECT}  |  {title}"
    p.style = doc.styles["TH Small"]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.runs[0].font.color.rgb = RGBColor.from_string(accent)
    p.runs[0].font.bold = True
    p_pr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "5")
    bottom.set(qn("w:color"), accent)
    borders.append(bottom)
    p_pr.append(borders)

    footer = section.footer
    p = footer.paragraphs[0]
    p.style = doc.styles["TH Small"]
    p.add_run("Tài liệu dự thảo – Chỉ sử dụng để xác nhận phạm vi  |  ")
    add_page_number(p)

    core = doc.core_properties
    core.title = title
    core.subject = "Phạm vi chức năng và lộ trình triển khai TimeHouse"
    core.author = VENDOR
    core.keywords = "TimeHouse, scope of work, phase, deliverables, acceptance criteria"


def add_cover(doc, doc_title, subtitle, accent, status, code):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run(PROJECT)
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(14)
    r.font.color.rgb = RGBColor.from_string(accent)

    doc.add_paragraph("")
    doc.add_paragraph("")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(doc_title.upper())
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(29)
    r.font.color.rgb = RGBColor.from_string(NAVY)

    p = doc.add_paragraph()
    r = p.add_run(subtitle)
    r.font.name = "Arial"
    r.font.size = Pt(15)
    r.font.color.rgb = RGBColor.from_string(accent)
    r.bold = True

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    r = p.add_run(PROJECT_LONG)
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor.from_string(DARK_GRAY)

    table = doc.add_table(rows=5, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.columns[0].width = Cm(5.0)
    table.columns[1].width = Cm(11.5)
    fields = [
        ("Mã tài liệu", code),
        ("Phiên bản", VERSION),
        ("Ngày phát hành", DATE),
        ("Trạng thái", status),
        ("Khách hàng", CLIENT),
    ]
    for row, (label, value) in zip(table.rows, fields):
        row.height = Cm(0.8)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        set_cant_split(row)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
        set_cell_shading(row.cells[0], accent)
        rp = row.cells[0].paragraphs[0]
        rr = rp.add_run(label)
        rr.bold = True
        rr.font.color.rgb = RGBColor(255, 255, 255)
        rr.font.size = Pt(9)
        vp = row.cells[1].paragraphs[0]
        vr = vp.add_run(value)
        vr.font.size = Pt(9)
        borders = {edge: {"val": "single", "sz": "5", "space": "0", "color": MID_GRAY} for edge in ("top", "start", "bottom", "end")}
        set_cell_border(row.cells[0], **borders)
        set_cell_border(row.cells[1], **borders)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(42)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(f"Đơn vị thực hiện: {VENDOR}\nTài liệu dành cho: {CLIENT}")
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor.from_string(DARK_GRAY)

    doc.add_page_break()


def add_document_control(doc, status, purpose):
    doc.add_heading("Kiểm soát tài liệu", level=1)
    add_table(
        doc,
        ["Phiên bản", "Ngày", "Trạng thái", "Mục đích phát hành"],
        [[VERSION, DATE, status, purpose]],
        [2.2, 2.7, 3.3, 8.0],
        NAVY,
    )
    doc.add_paragraph("")
    add_table(
        doc,
        ["Vai trò", "Đơn vị/Họ tên", "Trách nhiệm"],
        [
            ["Đại diện khách hàng", CLIENT, "Rà soát nghiệp vụ và xác nhận phạm vi"],
            ["Đơn vị thực hiện", VENDOR, "Phân tích, triển khai và bàn giao theo phạm vi được duyệt"],
        ],
        [4.0, 6.0, 6.2],
        NAVY,
    )


def add_toc(doc):
    doc.add_heading("Mục lục", level=1)
    p = doc.add_paragraph()
    r = p.add_run()
    add_field(r, 'TOC \\o "1-3" \\h \\z \\u')
    p2 = doc.add_paragraph("Mục lục sẽ tự cập nhật khi mở bằng Microsoft Word. Nếu chưa hiển thị, chọn toàn bộ tài liệu và nhấn F9.")
    p2.style = doc.styles["TH Small"]
    doc.add_page_break()


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    set_repeat_heading(p)
    return p


def add_bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)


def add_numbers(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)


def add_note(doc, title, text, accent, fill):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, 160, 180, 160, 180)
    set_cell_border(cell, start={"val": "single", "sz": "18", "space": "0", "color": accent})
    p = cell.paragraphs[0]
    r = p.add_run(title + "\n")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(accent)
    r2 = p.add_run(text)
    r2.font.color.rgb = RGBColor.from_string(DARK_GRAY)


def add_table(doc, headers, rows, widths, accent, font_size=8.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for idx, (cell, header) in enumerate(zip(hdr.cells, headers)):
        cell.width = Cm(widths[idx])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, accent)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_repeat_header_text(p.add_run(header))
    for ridx, values in enumerate(rows):
        row = table.add_row()
        set_cant_split(row)
        for idx, (cell, value) in enumerate(zip(row.cells, values)):
            cell.width = Cm(widths[idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            set_cell_margins(cell)
            if ridx % 2 == 1:
                set_cell_shading(cell, "F8FAFB")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(value))
            r.font.size = Pt(font_size)
            r.font.name = "Arial"
    border = {"val": "single", "sz": "4", "space": "0", "color": MID_GRAY}
    for row in table.rows:
        for cell in row.cells:
            set_cell_border(cell, top=border, start=border, bottom=border, end=border)
    doc.add_paragraph("")
    return table


def add_approval(doc):
    add_heading(doc, "Xác nhận phạm vi", level=1)
    doc.add_paragraph(
        "Việc ký xác nhận cho biết hai bên đã rà soát và thống nhất phạm vi ở mức nêu trong tài liệu. "
        "Thời gian, chi phí, SLA và điều khoản thương mại chỉ có hiệu lực khi được xác nhận trong kế hoạch hoặc phụ lục riêng."
    )
    add_table(
        doc,
        ["ĐẠI DIỆN KHÁCH HÀNG", "ĐẠI DIỆN ĐƠN VỊ THỰC HIỆN"],
        [[
            "Họ tên: ______________________________\nChức vụ: _____________________________\nNgày: ______ / ______ / __________\n\n\nChữ ký: ______________________________",
            "Họ tên: ______________________________\nChức vụ: _____________________________\nNgày: ______ / ______ / __________\n\n\nChữ ký: ______________________________",
        ]],
        [8.1, 8.1],
        NAVY,
        font_size=9,
    )


def add_common_governance(doc, phase_name, accent):
    add_heading(doc, "Giả định và phụ thuộc", level=1)
    add_bullets(doc, [
        "Khách hàng cung cấp dữ liệu mẫu, biểu mẫu, quy trình hiện hành và đầu mối có thẩm quyền xác nhận nghiệp vụ.",
        "Dữ liệu sử dụng cho kiểm thử và chuyển đổi phải đúng định dạng, đủ chất lượng và được khách hàng xác nhận quyền sử dụng.",
        "Tích hợp bên thứ ba phụ thuộc tài khoản, hợp đồng dịch vụ, API, hạn mức, chi phí và môi trường kiểm thử do nhà cung cấp hỗ trợ.",
        "Thời gian và chi phí chỉ được chốt sau khi phạm vi chi tiết, nguồn lực, dữ liệu và phụ thuộc tích hợp được xác nhận.",
        "Các nội dung kế toán, thuế, đầu tư và dữ liệu cá nhân cần được khách hàng hoặc đơn vị tư vấn chuyên môn xác nhận trước khi đưa vào production.",
    ])

    add_heading(doc, "Mốc triển khai dự kiến", level=1)
    add_table(
        doc,
        ["Mốc", "Nội dung", "Ngày dự kiến", "Điều kiện hoàn thành"],
        [
            ["M1", "Xác nhận phạm vi và yêu cầu chi tiết", "TBD", "Biên bản/tài liệu được hai bên xác nhận"],
            ["M2", "Thiết kế hoặc cập nhật mockup", "TBD", "Khách hàng duyệt luồng và giao diện chính"],
            ["M3", "Phát triển và kiểm thử nội bộ", "TBD", "Đạt tiêu chí kỹ thuật nội bộ"],
            ["M4", "UAT với khách hàng", "TBD", "Các lỗi mức nghiêm trọng đã được xử lý"],
            ["M5", "Bàn giao/Go-live", "TBD", "Biên bản nghiệm thu và kế hoạch vận hành được duyệt"],
        ],
        [1.3, 6.4, 3.0, 5.5],
        accent,
    )

    add_heading(doc, "Quy trình nghiệm thu", level=1)
    add_numbers(doc, [
        "Đơn vị thực hiện bàn giao phiên bản kiểm thử, tài liệu hướng dẫn và bộ dữ liệu/kịch bản kiểm thử đã thống nhất.",
        "Khách hàng thực hiện UAT theo tiêu chí nghiệm thu và ghi nhận lỗi hoặc sai khác so với phạm vi đã duyệt.",
        "Đơn vị thực hiện xử lý các lỗi thuộc phạm vi; các yêu cầu mới được đưa qua quy trình thay đổi phạm vi.",
        "Hai bên kiểm tra lại và lập biên bản xác nhận hoàn thành hoặc danh sách tồn đọng có thời hạn xử lý.",
    ])

    add_heading(doc, "Quản lý thay đổi phạm vi", level=1)
    doc.add_paragraph(
        f"Yêu cầu làm phát sinh chức năng, vai trò, báo cáo, tích hợp, luồng phê duyệt hoặc quy tắc tính toán chưa được mô tả trong {phase_name} "
        "được xem là thay đổi phạm vi. Mọi thay đổi cần được ghi nhận bằng Change Request, phân tích ảnh hưởng đến dữ liệu, bảo mật, thời gian và chi phí, "
        "sau đó được hai bên phê duyệt bằng văn bản trước khi triển khai."
    )


def make_overview():
    title = "Lộ trình và phạm vi chức năng theo 3 giai đoạn"
    doc = Document()
    configure_document(doc, title, ORANGE)
    add_cover(doc, "Phạm vi chức năng và lộ trình triển khai", "Tổng quan Phase 1 – Phase 2 – Phase 3", ORANGE, "Dự thảo gửi khách hàng xác nhận", "TH-SCOPE-OVERVIEW-1.0")
    add_document_control(doc, "Dự thảo", "Thống nhất cách phân chia phạm vi và định hướng triển khai")
    add_toc(doc)

    add_heading(doc, "1. Tóm tắt điều hành", 1)
    doc.add_paragraph(
        "TimeHouse được chia thành ba giai đoạn để ưu tiên khả năng vận hành cho thuê trước, sau đó mở rộng sang kinh doanh, tự động hóa và quản trị doanh nghiệp. "
        "Cách chia này giúp kiểm soát phạm vi, cho phép kiểm chứng giá trị sớm và hạn chế việc phát triển đồng thời quá nhiều chức năng chưa có dữ liệu nền."
    )
    add_note(
        doc,
        "HIỆN TRẠNG DỰ ÁN",
        "Code hiện tại thuộc Phase 1 và là mockup tương tác dùng để trình diễn, rà soát nghiệp vụ và chốt yêu cầu. Dữ liệu đang lưu trên trình duyệt; Zalo và các dịch vụ ngoài hệ thống đang được mô phỏng, chưa phải kết nối production.",
        BLUE,
        LIGHT_BLUE,
    )

    add_heading(doc, "2. Lộ trình ba giai đoạn", 1)
    add_table(
        doc,
        ["Giai đoạn", "Trọng tâm", "Giá trị mang lại", "Trạng thái"],
        [
            ["Phase 1 – Core Rental / Go-live", "Vận hành cho thuê cốt lõi", "Số hóa luồng từ phòng, khách, hợp đồng đến hóa đơn, thu tiền và trả phòng", "Hiện tại – đã có mockup tương tác"],
            ["Phase 2 – Sales, Automation & Operations", "Kinh doanh và tối ưu vận hành", "CRM, tự động hóa, bảo trì, đối soát và báo cáo quản trị", "Dự kiến – chưa triển khai"],
            ["Phase 3 – Enterprise & Investment", "Quản trị doanh nghiệp và đầu tư", "Tài sản, nhân sự, cổ đông, đầu tư và tài chính tích hợp", "Định hướng – chưa triển khai"],
        ],
        [4.3, 3.4, 6.0, 3.2],
        ORANGE,
    )
    add_note(doc, "LUỒNG LÕI PHASE 1", "Tòa nhà → Phòng → Khách thuê → Hợp đồng → Dịch vụ/Điện nước → Hóa đơn → Công nợ/Thu tiền → Zalo → Trả phòng/Hoàn cọc", BLUE, LIGHT_BLUE)

    add_heading(doc, "3. Phạm vi tóm tắt từng phase", 1)
    add_heading(doc, "3.1. Phase 1 – Core Rental / Go-live", 2)
    doc.add_paragraph("Mục tiêu: hình thành quy trình vận hành cho thuê tối thiểu trên một hệ thống thống nhất.")
    add_bullets(doc, [
        "Đăng nhập, tài khoản và phân quyền theo vai trò/phạm vi tòa.",
        "Tòa nhà, phòng, chủ nhà, khách thuê và hợp đồng thuê.",
        "Dịch vụ, bảng giá, chỉ số điện nước và lập hóa đơn theo kỳ.",
        "Công nợ, thu tiền, hoàn cọc, chi phí và báo cáo cơ bản.",
        "Thông báo hóa đơn/nhắc công nợ qua luồng Zalo; import dữ liệu cơ bản.",
    ])
    add_heading(doc, "3.2. Phase 2 – Sales, Automation & Operations", 2)
    doc.add_paragraph("Mục tiêu: tăng hiệu quả kinh doanh và giảm thao tác vận hành thủ công sau khi dữ liệu lõi ổn định.")
    add_bullets(doc, [
        "CRM/Lead, pipeline, lịch xem, giữ chỗ, chốt thuê và hoa hồng.",
        "OCR hợp đồng, bảo trì/bảo dưỡng và lịch công việc.",
        "Import bảng kê, đối soát, Data Job và xử lý ngoại lệ.",
        "Quản lý kỳ, dòng tiền, chi phí phân bổ và báo cáo quản trị.",
        "Rule Builder và thông báo đa sự kiện/kênh.",
    ])
    add_heading(doc, "3.3. Phase 3 – Enterprise & Investment", 2)
    doc.add_paragraph("Mục tiêu: mở rộng TimeHouse thành nền tảng quản trị doanh nghiệp và đầu tư.")
    add_bullets(doc, [
        "Tài sản, QR, kiểm kê, tình trạng và khấu hao.",
        "Nhân sự, phân công, KPI, chấm công và lương thưởng khi có nhu cầu.",
        "Cổ đông, vốn góp, phân phối lợi nhuận và cổng nhà đầu tư.",
        "QR Payment, cổng thanh toán, ngân hàng và đối soát tự động.",
        "Audit nâng cao, BI, báo cáo tùy chỉnh và cảnh báo bất thường.",
    ])

    add_heading(doc, "4. Ma trận phân bổ chức năng", 1)
    add_table(
        doc,
        ["Nhóm chức năng", "Phase 1", "Phase 2", "Phase 3"],
        [
            ["Tòa nhà, phòng, khách thuê", "Phạm vi chính", "Mở rộng vận hành", "Phân tích nâng cao"],
            ["Hợp đồng và dịch vụ", "Phạm vi chính", "OCR/automation", "–"],
            ["Hóa đơn, công nợ, thu tiền", "Cơ bản", "Đối soát/quản lý kỳ", "Ngân hàng/thanh toán"],
            ["Trả phòng và hoàn cọc", "Cơ bản", "Workflow nâng cao", "–"],
            ["Zalo/thông báo", "Cơ bản", "Rule đa sự kiện/kênh", "–"],
            ["Báo cáo", "Vận hành cơ bản", "Quản trị", "BI/tùy chỉnh"],
            ["CRM, sale, hoa hồng", "–", "Phạm vi chính", "–"],
            ["Bảo trì/bảo dưỡng", "–", "Phạm vi chính", "Liên kết tài sản"],
            ["Tài sản/kiểm kê", "Danh mục cơ bản", "Theo dõi bảo trì", "Phạm vi chính"],
            ["Nhân sự", "–", "–", "Phạm vi chính"],
            ["Cổ đông và đầu tư", "–", "–", "Phạm vi chính"],
        ],
        [6.0, 3.4, 3.9, 3.8],
        ORANGE,
    )

    add_heading(doc, "5. Nguyên tắc bàn giao và nghiệm thu", 1)
    add_bullets(doc, [
        "Mỗi phase có tài liệu phạm vi riêng, danh sách deliverable và tiêu chí nghiệm thu tương ứng.",
        "Một phase chỉ được xác nhận hoàn thành khi các luồng nghiệp vụ trọng tâm chạy xuyên suốt trên dữ liệu kiểm thử đã thống nhất.",
        "Yêu cầu production phải được kiểm thử thêm về phân quyền, bảo mật, hiệu năng, sao lưu, dữ liệu thật và tích hợp thật.",
        "Hạng mục chưa có công thức hoặc quy trình được duyệt sẽ giữ trạng thái TBD và không được hiểu là đã bao gồm trong cam kết triển khai.",
    ])

    add_heading(doc, "6. Tài liệu bàn giao theo phase", 1)
    add_table(
        doc,
        ["Mã", "Tài liệu", "Mục đích"],
        [
            ["TH-SOW-P1", "Phạm vi Phase 1 – Core Rental / Go-live", "Xác nhận phạm vi hiện tại và tiêu chí Go-live"],
            ["TH-SOW-P2", "Phạm vi Phase 2 – Sales, Automation & Operations", "Chốt định hướng CRM, automation và vận hành nâng cao"],
            ["TH-SOW-P3", "Phạm vi Phase 3 – Enterprise & Investment", "Chốt định hướng doanh nghiệp, tài sản và đầu tư"],
        ],
        [2.8, 7.0, 7.3],
        ORANGE,
    )

    add_common_governance(doc, "lộ trình ba phase", ORANGE)
    add_approval(doc)
    path = OUT / "00_TimeHouse_Tong-quan-lo-trinh_3-Phase_v1.0.docx"
    doc.save(path)
    return path


PHASES = {
    1: {
        "accent": BLUE,
        "fill": LIGHT_BLUE,
        "name": "Phase 1 – Core Rental / Go-live",
        "code": "TH-SOW-P1-1.0",
        "status": "Hiện tại – đã có mockup tương tác",
        "purpose": "Xác nhận phạm vi hiện tại và tiêu chí triển khai production/Go-live",
        "objective": "Hoàn thiện luồng nghiệp vụ tối thiểu để doanh nghiệp quản lý hoạt động cho thuê hằng ngày trên một hệ thống thống nhất, giảm theo dõi rời rạc và tạo nền tảng dữ liệu cho các phase tiếp theo.",
        "scope": [
            ["P1-F01", "Đăng nhập và phân quyền", "Đăng nhập/đăng xuất; tài khoản; vai trò Admin, Kế toán, Quản lý/Vận hành; phạm vi theo tòa", "Người dùng chỉ xem và thao tác đúng quyền"],
            ["P1-F02", "Tổng quan", "Tình trạng phòng; phải thu, đã thu, còn nợ, quá hạn; việc cần xử lý", "Nắm nhanh tình hình vận hành và tài chính"],
            ["P1-F03", "Tòa nhà và chủ nhà", "Hồ sơ tòa, người phụ trách, chủ nhà, hợp đồng đầu vào và lịch thanh toán cơ bản", "Quản lý tập trung thông tin đầu vào"],
            ["P1-F04", "Phòng", "Danh sách/chi tiết; giá; dịch vụ; import; trạng thái Sẵn sàng, Đang thuê, Chờ dọn, Bảo trì, Ngừng sử dụng", "Theo dõi khả năng khai thác từng phòng"],
            ["P1-F05", "Khách thuê", "Hồ sơ, liên hệ/Zalo, giấy tờ, phương tiện, lịch sử thuê, công nợ và import", "Hồ sơ khách liên kết với phòng/hợp đồng"],
            ["P1-F06", "Hợp đồng thuê", "Tạo nháp, kích hoạt, gia hạn, kết thúc/hủy; giá, cọc, kỳ, thành viên, dịch vụ và tệp đính kèm", "Quản lý vòng đời hợp đồng và trạng thái phòng"],
            ["P1-F07", "Dịch vụ và điện nước", "Danh mục/bảng giá; nhập/import chỉ số; đối chiếu kỳ trước; cảnh báo bất thường", "Chuẩn hóa dữ liệu đầu vào để tính tiền"],
            ["P1-F08", "Hóa đơn", "Tạo đơn lẻ/hàng loạt; nháp, phát hành, hủy/điều chỉnh; chi tiết phí và bản in/PDF", "Lập hóa đơn và theo dõi trạng thái thanh toán"],
            ["P1-F09", "Công nợ và thu tiền", "Phải thu/đã thu/còn nợ/quá hạn; thu đủ/một phần; phân bổ; minh chứng; hoàn tác theo quyền", "Công nợ cập nhật theo khoản thu thực tế"],
            ["P1-F10", "Trả phòng và hoàn cọc", "Kết thúc HĐ; kiểm tra nợ; khấu trừ; duyệt; ghi nhận hoàn; Chờ dọn và xác nhận dọn xong", "Khép kín vòng đời thuê và đưa phòng về Sẵn sàng"],
            ["P1-F11", "Chi phí cơ bản", "Khoản chi theo nhóm, tòa, ngày, số tiền, chứng từ và ghi chú", "Có dữ liệu chi phí vận hành ban đầu"],
            ["P1-F12", "Thông báo Zalo", "Cấu hình mẫu/sự kiện; chọn người nhận; re-check công nợ; theo dõi kết quả; retry", "Hỗ trợ gửi hóa đơn và nhắc công nợ có kiểm soát"],
            ["P1-F13", "Báo cáo và import", "Báo cáo phòng/công nợ/thu tiền; import phòng, khách, chỉ số và hóa đơn", "Hỗ trợ vận hành, đối chiếu và cập nhật hàng loạt"],
        ],
        "flow": [
            "Tạo hoặc import khách thuê.",
            "Chọn phòng Sẵn sàng; tạo và kích hoạt hợp đồng với giá, cọc, kỳ thanh toán và dịch vụ.",
            "Ghi nhận chỉ số điện nước; tạo và phát hành hóa đơn theo kỳ.",
            "Gửi thông báo hóa đơn; xem công nợ và ghi nhận thu đủ hoặc một phần.",
            "Gửi nhắc công nợ; theo dõi trạng thái và thử lại trường hợp gửi lỗi.",
            "Kết thúc hợp đồng; lập, duyệt và ghi nhận hồ sơ hoàn cọc.",
            "Chuyển phòng sang Chờ dọn; xác nhận dọn xong và đưa phòng về Sẵn sàng.",
            "Kiểm tra báo cáo phòng, công nợ và thu tiền theo kỳ.",
        ],
        "deliverables": [
            ["P1-D01", "Bộ màn hình và luồng Phase 1", "Các nhóm chức năng P1-F01 đến P1-F13", "Các màn/luồng được duyệt theo yêu cầu chi tiết", "Đang có mockup"],
            ["P1-D02", "Dữ liệu và import", "File mẫu, kiểm tra dữ liệu, kết quả import", "Import thành công dữ liệu hợp lệ; nêu rõ dòng lỗi", "Đang có mockup"],
            ["P1-D03", "Phân quyền", "Ba vai trò Phase 1 và phạm vi theo tòa", "Kịch bản quyền được kiểm thử và không lộ dữ liệu ngoài phạm vi", "Đang có mockup"],
            ["P1-D04", "Tích hợp thông báo", "Kết nối Zalo production sau khi có tài khoản/API", "Gửi/nhận kết quả qua môi trường thật và có log đối soát", "Chưa tích hợp thật"],
            ["P1-D05", "Báo cáo và xuất dữ liệu", "Phòng, công nợ, thu tiền", "Số liệu đối chiếu đúng với bộ dữ liệu chuẩn", "Đang có mockup"],
            ["P1-D06", "Tài liệu bàn giao", "Hướng dẫn sử dụng, cấu hình, triển khai và biên bản UAT", "Tài liệu được bàn giao và khách hàng xác nhận", "Chờ production"],
        ],
        "out": [
            "CRM/Lead, lịch xem phòng, giữ chỗ tự động, quản lý kênh và hoa hồng sale.",
            "OCR/AI trích xuất hợp đồng, chữ ký điện tử và phê duyệt hợp đồng nhiều cấp.",
            "Import bảng kê ngân hàng, đối soát hoặc tự động khớp giao dịch.",
            "Bảo trì/sửa chữa, lịch bảo dưỡng và quản lý nhà cung cấp nâng cao.",
            "Workflow hoàn cọc nhiều cấp hoặc theo ngưỡng số tiền.",
            "Rule Builder đa kênh, fallback SMS/email và hội thoại hai chiều.",
            "BI, lợi nhuận, dòng tiền, khóa kỳ và trình tạo báo cáo tùy chỉnh.",
            "Nhân sự, cổ đông, đầu tư và cổng nhà đầu tư.",
        ],
        "acceptance": [
            "Hoàn thành xuyên suốt luồng từ tạo khách, hợp đồng, hóa đơn, thu tiền đến trả phòng trên dữ liệu UAT.",
            "Trạng thái phòng, hợp đồng, hóa đơn, công nợ và hoàn cọc chuyển đúng theo quy tắc đã chốt.",
            "Ba vai trò Phase 1 truy cập đúng chức năng và dữ liệu được phân công.",
            "Số liệu báo cáo đối chiếu đúng với dữ liệu nguồn và công thức được duyệt.",
            "Các tích hợp production có log kết quả và cơ chế xử lý lỗi theo tài liệu tích hợp.",
            "Không còn lỗi mức Blocker/Critical; các lỗi còn lại có phương án và thời hạn được khách hàng chấp thuận.",
        ],
    },
    2: {
        "accent": GREEN,
        "fill": LIGHT_GREEN,
        "name": "Phase 2 – Sales, Automation & Operations",
        "code": "TH-SOW-P2-1.0",
        "status": "Dự kiến – chưa triển khai",
        "purpose": "Xác nhận định hướng CRM, tự động hóa và vận hành nâng cao",
        "objective": "Tối ưu hoạt động kinh doanh và vận hành sau khi dữ liệu lõi Phase 1 ổn định; giảm thao tác thủ công, tăng khả năng kiểm soát và cung cấp báo cáo quản trị sâu hơn.",
        "scope": [
            ["P2-F01", "CRM và Lead", "Danh sách/Kanban Lead; nguồn; sale phụ trách; lịch sử chăm sóc; pipeline", "Theo dõi cơ hội và tỷ lệ chuyển đổi"],
            ["P2-F02", "Lịch xem và giữ chỗ", "Đặt/nhắc lịch; kết quả xem; phí/thời hạn giữ; tự hết hạn; khóa phòng", "Giảm xung đột và chuẩn hóa quy trình trước hợp đồng"],
            ["P2-F03", "Chốt thuê và hoa hồng", "Chuyển Lead thành khách/HĐ; giao dịch, doanh số và điều kiện hoa hồng", "Liên kết hiệu suất sale với dữ liệu vận hành"],
            ["P2-F04", "OCR hợp đồng", "Tải PDF/JPG; trích xuất; confidence; rà soát; mapping vào form", "Giảm thời gian nhập liệu nhưng vẫn có bước kiểm tra"],
            ["P2-F05", "Bảo trì/bảo dưỡng", "Sự cố; ưu tiên; phân công; tiến độ; lịch; nhà cung cấp; checklist; bằng chứng", "Theo dõi trách nhiệm, SLA và chi phí bảo trì"],
            ["P2-F06", "Import và đối soát", "Bảng kê thu tiền; mapping; auto-match; ngoại lệ; Data Job; retry dòng lỗi", "Giảm nhập liệu và đối soát thủ công"],
            ["P2-F07", "Tài chính nâng cao", "Chi phí phân bổ; sổ cọc; thu chi; dòng tiền; kết quả vận hành; khóa kỳ", "Kiểm soát tài chính theo tòa và kỳ"],
            ["P2-F08", "Thông báo nâng cao", "Rule Builder; nhiều sự kiện; phiên bản mẫu; lịch gửi; provider response; fallback", "Tự động hóa thông báo có cấu hình và truy vết"],
            ["P2-F09", "Báo cáo quản trị", "Lấp đầy; thời gian trống; lợi nhuận; sale; marketing; bảo trì; Report Hub", "Cung cấp dữ liệu quản trị và ra quyết định"],
        ],
        "flow": [
            "Tiếp nhận Lead, phân công sale và quản lý pipeline.",
            "Đặt lịch xem, ghi nhận kết quả và giữ chỗ phòng trong thời hạn cho phép.",
            "Chốt thuê, chuyển dữ liệu sang khách thuê/hợp đồng và ghi nhận hoa hồng.",
            "Tải hợp đồng, OCR, rà soát trường và cập nhật form hợp đồng.",
            "Nhập bảng kê, tự khớp giao dịch với hóa đơn và xử lý ngoại lệ.",
            "Tiếp nhận sự cố, phân công, theo dõi xử lý và ghi nhận chi phí/bằng chứng.",
            "Khóa kỳ và xem các báo cáo quản trị đã đối chiếu.",
        ],
        "deliverables": [
            ["P2-D01", "CRM và pipeline", "Lead, lịch xem, giữ chỗ, chốt thuê", "Luồng Lead → Hợp đồng chạy xuyên suốt", "Chưa triển khai"],
            ["P2-D02", "Hoa hồng và hiệu suất sale", "Quy tắc, giao dịch, báo cáo", "Kết quả khớp bộ số liệu chuẩn", "Chưa triển khai"],
            ["P2-D03", "OCR hợp đồng", "Trích xuất và màn review", "Người dùng sửa/xác nhận trước khi lưu; đo độ chính xác theo mẫu", "Chưa triển khai"],
            ["P2-D04", "Bảo trì/bảo dưỡng", "Sự cố, lịch, phân công, checklist", "Luồng từ tiếp nhận đến hoàn tất có lịch sử và bằng chứng", "Chưa triển khai"],
            ["P2-D05", "Đối soát và Data Job", "Import, auto-match, ngoại lệ, retry", "Không tạo trùng; truy vết được từng dòng xử lý", "Chưa triển khai"],
            ["P2-D06", "Báo cáo quản trị", "Report Hub và chỉ số Phase 2", "Công thức và số liệu được khách hàng xác nhận", "Chưa triển khai"],
        ],
        "out": [
            "Quản lý tài sản bằng QR, kiểm kê và khấu hao toàn diện.",
            "Nhân sự, chấm công, KPI và bảng lương.",
            "Cổ đông, vốn góp, phân phối lợi nhuận và cổng nhà đầu tư.",
            "QR Payment, cổng thanh toán, kết nối ngân hàng và webhook.",
            "Kho dữ liệu, BI doanh nghiệp và phát hiện bất thường nâng cao.",
            "Các chức năng hoặc tích hợp chưa có quy trình/công thức được phê duyệt.",
        ],
        "acceptance": [
            "Luồng Lead → lịch xem → giữ chỗ → chốt thuê → hợp đồng chạy xuyên suốt trên dữ liệu UAT.",
            "Quy tắc hoa hồng và chỉ số sale khớp bộ số liệu chuẩn được khách hàng duyệt.",
            "OCR hiển thị độ tin cậy, cho phép rà soát và không tự ghi dữ liệu chưa được xác nhận.",
            "Luồng sự cố/bảo trì lưu được phân công, lịch sử trạng thái, chi phí và bằng chứng.",
            "Đối soát không tạo trùng, có danh sách ngoại lệ và truy vết được từng dòng.",
            "Báo cáo Phase 2 đối chiếu đúng với công thức và dữ liệu nguồn đã chốt.",
        ],
    },
    3: {
        "accent": PURPLE,
        "fill": LIGHT_PURPLE,
        "name": "Phase 3 – Enterprise & Investment",
        "code": "TH-SOW-P3-1.0",
        "status": "Định hướng – chưa triển khai",
        "purpose": "Xác nhận định hướng quản trị doanh nghiệp, tài sản và đầu tư",
        "objective": "Mở rộng TimeHouse từ hệ thống vận hành cho thuê thành nền tảng quản trị doanh nghiệp và đầu tư trên cơ sở dữ liệu đã hình thành ở Phase 1–2.",
        "scope": [
            ["P3-F01", "Quản lý tài sản", "Tài sản theo phòng/tòa; QR; kiểm kê; tình trạng; mất/hỏng; khấu hao; lịch sử", "Kiểm soát vòng đời và chênh lệch kiểm kê"],
            ["P3-F02", "Nhân sự", "Hồ sơ; phòng ban/chức danh; phân công; KPI; chấm công; bảng lương khi cần", "Liên kết nguồn lực với hoạt động vận hành"],
            ["P3-F03", "Cổ đông và vốn góp", "Cổ đông; tỷ lệ; đợt góp; đã góp/còn thiếu; lịch sử", "Minh bạch tình trạng vốn góp"],
            ["P3-F04", "Phân phối lợi nhuận", "Tỷ lệ; kỳ phân phối; phải trả/đã trả; hiệu quả đầu tư", "Theo dõi nghĩa vụ và kết quả đầu tư"],
            ["P3-F05", "Cổng nhà đầu tư", "Dashboard vốn, lợi nhuận, lịch phân phối và tài liệu", "Kênh tra cứu riêng cho nhà đầu tư"],
            ["P3-F06", "Thanh toán và ngân hàng", "QR Payment; gateway; ngân hàng; webhook; đối soát; auto-match", "Tự động hóa thu tiền và đối soát"],
            ["P3-F07", "Audit và BI", "Nhật ký nâng cao; dashboard BI; báo cáo tùy chỉnh; cảnh báo/anomaly", "Tăng khả năng kiểm soát và phân tích"],
        ],
        "flow": [
            "Gắn tài sản với phòng/tòa, tạo mã QR và thực hiện kiểm kê định kỳ.",
            "Ghi nhận nhân sự, phân công và dữ liệu KPI/chấm công theo chính sách đã duyệt.",
            "Tạo dự án/cơ cấu cổ đông, ghi nhận đợt góp vốn và số dư còn thiếu.",
            "Tính và duyệt kỳ phân phối lợi nhuận; công bố cho nhà đầu tư.",
            "Tiếp nhận thanh toán qua ngân hàng/cổng thanh toán và tự động đối soát.",
            "Theo dõi audit, dashboard BI và cảnh báo bất thường.",
        ],
        "deliverables": [
            ["P3-D01", "Tài sản và kiểm kê", "Danh mục, QR, kiểm kê, khấu hao", "Kiểm kê xác định đủ chênh lệch và lịch sử", "Chưa triển khai"],
            ["P3-D02", "Nhân sự", "Hồ sơ, phân công, KPI/chấm công/lương", "Kết quả đúng chính sách được phê duyệt", "Chưa triển khai"],
            ["P3-D03", "Cổ đông và vốn góp", "Cơ cấu, đợt góp, số dư", "Đối chiếu đúng với bộ số liệu pháp lý/tài chính", "Chưa triển khai"],
            ["P3-D04", "Phân phối và Investor Portal", "Kỳ phân phối, dashboard, tài liệu", "Nhà đầu tư chỉ xem đúng dữ liệu được cấp quyền", "Chưa triển khai"],
            ["P3-D05", "Thanh toán/ngân hàng", "API, webhook, auto-match, đối soát", "Giao dịch được xác thực, chống trùng và truy vết", "Chưa triển khai"],
            ["P3-D06", "Audit và BI", "Nhật ký, dashboard, cảnh báo", "Dữ liệu đúng nguồn, quyền truy cập và thời gian lưu", "Chưa triển khai"],
        ],
        "out": [
            "Tư vấn pháp lý, thuế, chứng khoán hoặc quản lý quỹ.",
            "Thay thế phần mềm kế toán hoặc ERP nếu chưa có thỏa thuận tích hợp riêng.",
            "Mua tài khoản, phí API, thiết bị QR, thiết bị chấm công hoặc dịch vụ ngân hàng/cổng thanh toán.",
            "Kho dữ liệu hoặc mô hình AI chưa có nguồn dữ liệu, tiêu chí chất lượng và chính sách quản trị được duyệt.",
            "Chức năng đa quốc gia, đa chuẩn kế toán hoặc đa tiền tệ nếu chưa được mô tả trong phụ lục.",
        ],
        "acceptance": [
            "Số lượng và tình trạng tài sản sau kiểm kê đối chiếu đúng với dữ liệu chuẩn; mọi điều chỉnh có lịch sử.",
            "KPI, chấm công và lương thưởng khớp chính sách/công thức đã được khách hàng phê duyệt.",
            "Vốn góp và phân phối lợi nhuận đối chiếu đúng bộ số liệu pháp lý/tài chính được cung cấp.",
            "Cổng nhà đầu tư bảo đảm phân quyền, riêng tư và chỉ hiển thị đúng dữ liệu của người dùng.",
            "Tích hợp thanh toán/ngân hàng có xác thực, chống ghi nhận trùng, nhật ký và cơ chế xử lý ngoại lệ.",
            "Dashboard BI và cảnh báo truy xuất được nguồn dữ liệu, công thức và quyền truy cập.",
        ],
    },
}


def make_phase(number):
    d = PHASES[number]
    doc = Document()
    configure_document(doc, d["name"], d["accent"])
    add_cover(doc, f"Phạm vi công việc – {d['name']}", "Software Scope of Work", d["accent"], d["status"], d["code"])
    add_document_control(doc, d["status"], d["purpose"])
    add_toc(doc)

    add_heading(doc, "1. Bối cảnh và mục tiêu", 1)
    doc.add_paragraph(d["objective"])
    if number == 1:
        add_note(
            doc,
            "TRẠNG THÁI HIỆN TẠI",
            "Phase 1 đã có mockup tương tác để demo và xác nhận nghiệp vụ. Mockup dùng localStorage trên trình duyệt; gửi Zalo và các dịch vụ ngoài hệ thống đang được mô phỏng. Việc triển khai production cần backend, cơ sở dữ liệu, hạ tầng, bảo mật, sao lưu và tích hợp thật theo kế hoạch riêng.",
            d["accent"],
            d["fill"],
        )
    else:
        add_note(doc, "TRẠNG THÁI PHẠM VI", f"{d['name']} là phạm vi dự kiến. Chức năng, công thức, tích hợp, thời gian và chi phí chỉ được cam kết sau bước khảo sát và xác nhận yêu cầu chi tiết.", d["accent"], d["fill"])

    add_heading(doc, "2. Phạm vi chức năng", 1)
    add_table(doc, ["Mã", "Nhóm chức năng", "Nội dung chính", "Kết quả mong đợi"], d["scope"], [1.6, 3.8, 7.3, 5.0], d["accent"], font_size=8)

    add_heading(doc, "3. Luồng nghiệp vụ trọng tâm", 1)
    add_numbers(doc, d["flow"])

    add_heading(doc, "4. Kết quả bàn giao (Deliverables)", 1)
    add_table(doc, ["Mã", "Kết quả bàn giao", "Nội dung", "Tiêu chí chấp nhận", "Hiện trạng"], d["deliverables"], [1.7, 3.5, 4.5, 5.5, 3.0], d["accent"], font_size=7.8)

    add_heading(doc, "5. Tiêu chí nghiệm thu", 1)
    add_bullets(doc, d["acceptance"])
    doc.add_paragraph(
        "Bộ dữ liệu UAT, mức độ lỗi, thời gian phản hồi và danh sách trình duyệt/thiết bị hỗ trợ sẽ được chốt trong Kế hoạch kiểm thử trước khi bắt đầu UAT."
    )

    add_heading(doc, "6. Ngoài phạm vi", 1)
    add_bullets(doc, d["out"])
    add_note(doc, "NGUYÊN TẮC PHẠM VI", "Nội dung không được nêu trong mục Phạm vi chức năng hoặc Kết quả bàn giao không mặc nhiên được hiểu là đã bao gồm.", d["accent"], d["fill"])

    add_heading(doc, "7. Trách nhiệm các bên", 1)
    add_table(
        doc,
        ["Bên", "Trách nhiệm chính"],
        [
            ["Khách hàng", "Cung cấp dữ liệu/quy trình; chỉ định người duyệt; phản hồi đúng hạn; cung cấp tài khoản dịch vụ bên thứ ba; tổ chức UAT và xác nhận kết quả."],
            ["Đơn vị thực hiện", "Phân tích yêu cầu; thiết kế/phát triển; kiểm thử nội bộ; sửa lỗi thuộc phạm vi; bàn giao tài liệu và hỗ trợ UAT theo kế hoạch."],
            ["Nhà cung cấp bên thứ ba", "Cung cấp API, tài khoản, môi trường, tài liệu kỹ thuật và hỗ trợ sự cố theo điều khoản dịch vụ riêng."],
        ],
        [4.0, 12.2],
        d["accent"],
    )

    add_common_governance(doc, d["name"], d["accent"])
    add_approval(doc)
    slug = {1: "01_TimeHouse_Phase-1_Core-Rental-Go-live_v1.0.docx", 2: "02_TimeHouse_Phase-2_Sales-Automation-Operations_v1.0.docx", 3: "03_TimeHouse_Phase-3_Enterprise-Investment_v1.0.docx"}[number]
    path = OUT / slug
    doc.save(path)
    return path


def validate_docx(path):
    required = {"[Content_Types].xml", "_rels/.rels", "word/document.xml", "word/styles.xml"}
    with ZipFile(path) as zf:
        names = set(zf.namelist())
        missing = required - names
        if missing:
            raise RuntimeError(f"{path.name}: missing DOCX parts: {sorted(missing)}")
        for name in required:
            if not zf.read(name):
                raise RuntimeError(f"{path.name}: empty DOCX part: {name}")
    # Reopen through python-docx to verify package relationships and XML parsing.
    check = Document(path)
    if not check.paragraphs or not any(p.text.strip() for p in check.paragraphs):
        raise RuntimeError(f"{path.name}: no readable paragraphs")


def package(paths):
    zip_path = OUT / "TimeHouse_Bo-tai-lieu-pham-vi_3-Phase_v1.0.zip"
    with ZipFile(zip_path, "w", ZIP_DEFLATED) as zf:
        for path in paths:
            zf.write(path, arcname=path.name)
    return zip_path


if __name__ == "__main__":
    generated = [make_overview(), make_phase(1), make_phase(2), make_phase(3)]
    for item in generated:
        validate_docx(item)
    bundle = package(generated)
    print("Generated:")
    for item in generated + [bundle]:
        print(f"- {item.relative_to(ROOT)} ({item.stat().st_size:,} bytes)")
