"""Sinh 2 file demo mà màn "Dữ liệu mẫu tháng 9" trỏ tới, từ trạng thái do generate-september-flow.mjs ghi ra.

    python -X utf8 scripts/seed/build_september_artifacts.py

Đầu vào : tmp/september-flow/state.json, docs/data/September_2026_Flow_Manifest.json
Đầu ra  : mockup/demo/september/contract-flow.pdf          (hợp đồng ca khách mới, bản mẫu)
          mockup/demo/september/Bo_du_lieu_mau_2026-09.xlsx (mỗi luồng một sheet, có cột nguồn)
Cả hai file sinh lặp lại cho cùng byte (bỏ dấu thời gian) để chạy lại không tạo diff.
"""
from datetime import datetime
from pathlib import Path
import io
import json
import os
import re
import sys
import zipfile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab import rl_config

rl_config.invariant = 1  # PDF không mang ngày tạo / ID ngẫu nhiên

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'scripts'))
import generate_matched_contracts as gc  # noqa: E402  (dùng lại font, kiểu chữ, bảng)
from reportlab.lib import colors  # noqa: E402
from reportlab.lib.pagesizes import A4  # noqa: E402
from reportlab.lib.units import mm  # noqa: E402
from reportlab.platypus import BaseDocTemplate, Frame, PageTemplate  # noqa: E402

OUT = ROOT / 'mockup' / 'demo' / 'september'
PERIOD = '2026-09'
FIXED = datetime(2026, 10, 5, 0, 0, 0)


def vnd(n):
    return f"{round(n or 0):,}".replace(',', '.')


def dmy(iso):
    return f"{iso[8:10]}/{iso[5:7]}/{iso[0:4]}" if iso else ''


def ref(rec):
    p = rec.get('provenance') or {}
    return p.get('sourceRef') or rec.get('evidence') or (rec.get('excel') or {}).get('src') or ''


# ---------- PDF hợp đồng ca mẫu ----------
def footer(canvas, doc):
    canvas.saveState()
    canvas.setFont('TimesVN', 8)
    canvas.setFillColor(colors.HexColor('#8a3d24'))
    canvas.drawString(18 * mm, 15 * mm, 'DỮ LIỆU MẪU THÁNG 9/2026 · KHÔNG CÓ GIÁ TRỊ GIAO KẾT')
    canvas.setFillColor(colors.black)
    canvas.drawRightString(192 * mm, 15 * mm, f'Trang {doc.page}')
    canvas.restoreState()


def contract_pdf(state, journey, path):
    by = {c: {r['id']: r for r in state[c]} for c in ('stays', 'rooms', 'buildings', 'customers', 'deals')}
    stay, room = by['stays'][journey['stayId']], by['rooms'][journey['roomId']]
    building, customer, deal = by['buildings'][journey['buildingId']], by['customers'][stay['customerId']], by['deals'][journey['dealId']]
    items = journey['rate'].get('items') or {}
    labels = {'electric': 'Điện', 'water': 'Nước', 'internet': 'Internet', 'combo': 'Dịch vụ chung', 'cleaning': 'Vệ sinh', 'elevator': 'Thang máy'}
    units = {'meter': 'theo chỉ số', 'person': '/người/tháng', 'room': '/phòng/tháng', 'vehicle': '/xe/tháng', 'fixed': '/tháng'}
    s = []
    gc.title(s, 'HỢP ĐỒNG CHO THUÊ PHÒNG', f"HĐKH-{building['code']}-{room['code']}-2026-09")
    gc.para(s, 'Ngày lập:', f"{dmy(deal['closeDate'])}. Địa điểm: tòa {building['code']}, {building.get('address') or ''}.")
    gc.para(s, 'BÊN A – BÊN CHO THUÊ PHÒNG:', 'Đơn vị khai thác tòa nhà theo hợp đồng thuê với chủ nhà (bản mẫu, không ghi thông tin cá nhân).')
    gc.para(s, 'BÊN B – KHÁCH THUÊ:', f"{customer['name']} (khách mẫu của bộ dữ liệu, không phải người thật).")
    gc.section(s, 'Điều 1. Phòng thuê',
               f"Bên A cho Bên B thuê <b>phòng {room['code']}</b> tại tòa <b>{building['code']}</b>. Số người ở: <b>01</b>.")
    gc.section(s, 'Điều 2. Thời hạn',
               f"Ngày nhận phòng và bắt đầu tính tiền: <b>{dmy(stay['rentStart'])}</b>. Thời hạn <b>{deal['term']} tháng</b>, đến hết <b>{dmy(stay['endDate'])}</b>.")
    gc.section(s, 'Điều 3. Giá thuê và tiền cọc',
               f"Tiền thuê: <b>{vnd(journey['rent'])} đồng/tháng</b>. Tiền cọc: <b>{vnd(journey['deposit'])} đồng</b>, đã nộp ngày {dmy(deal['closeDate'])}.",
               f"Hóa đơn đầu kỳ 09/2026 tính từ ngày nhận phòng: <b>{vnd(journey['invoiceTotal'])} đồng</b>, đã thu đủ.")
    rows = [['Dịch vụ', 'Đơn giá', 'Cách tính']]
    for key, it in items.items():
        rows.append([labels.get(key, key), f"{vnd(it.get('unit'))} đồng", units.get(it.get('method'), it.get('method') or '') + (f" × {it['qty']}" if it.get('qty') else '')])
    gc.section(s, 'Điều 4. Biểu phí dịch vụ',
               f"Biểu phí theo hợp đồng mẫu cùng tòa ({journey.get('rateFrom') or 'SRC-08'}). Chỉ số điện bàn giao: <b>{journey['meterStart']['el']}</b> kWh, nối tiếp chỉ số cuối của phòng.")
    s.append(gc.table(rows, [48 * mm, 54 * mm, 72 * mm]))
    gc.section(s, 'Ghi chú bộ dữ liệu',
               'Trong bộ dữ liệu demo, file này được đánh dấu "đã ký" để chạy luồng đủ điều kiện chi hoa hồng (CH-19). Bản PDF không mang chữ ký thật.')
    gc.signatures(s, 'BÊN A – BÊN CHO THUÊ PHÒNG', 'BÊN B – KHÁCH THUÊ')
    doc = BaseDocTemplate(str(path), pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=20 * mm, bottomMargin=23 * mm,
                          title='Hop-dong-mau-luong-thang-9', author='TimoHouse demo')
    doc.addPageTemplates(PageTemplate(id='contract', frames=Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, 0, 0, 0, 0), onPage=footer))
    doc.build(s)


# ---------- Excel bộ dữ liệu ----------
HEAD = Font(bold=True, color='FFFFFF')
FILL = PatternFill('solid', fgColor='1F4E79')


def sheet(wb, name, header, rows, money=()):
    ws = wb.create_sheet(name)
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill, c.alignment = HEAD, FILL, Alignment(wrap_text=True, vertical='top')
    for r in rows:
        ws.append(list(r))
    for i, h in enumerate(header, 1):
        width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:300] if r[i - 1] is not None]) + 2
        ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, width))
        if h in money:
            for cell in ws.iter_cols(min_col=i, max_col=i, min_row=2):
                for c in cell:
                    c.number_format = '#,##0.##'
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    return ws


def workbook(state, manifest, path):
    by = lambda c: {r['id']: r for r in state.get(c, [])}  # noqa: E731
    rooms, buildings, stays, emps, deals, shs = by('rooms'), by('buildings'), by('stays'), by('employees'), by('deals'), by('shareholders')
    code = lambda m, i, k='code': (m.get(i) or {}).get(k, '')  # noqa: E731
    paid = {}
    for p in state['payments']:
        for a in p.get('allocations') or []:
            paid[a['invoiceId']] = paid.get(a['invoiceId'], 0) + a['amount']
    wb = Workbook()
    ov = wb.active
    ov.title = 'Tổng quan'
    ov.append([manifest['label'], f"kỳ {PERIOD}", f"ngày xem {manifest['today']}", f"phiên bản {manifest['version']}"])
    ov['A1'].font = Font(bold=True, size=13)
    ov.append([])
    ov.append(['Bước', 'Kết quả chính'])
    for st in manifest['summary']['steps']:
        brief = {k: v for k, v in st.items() if k != 'name' and not isinstance(v, (dict, list))}
        ov.append([st['name'], json.dumps(brief, ensure_ascii=False)])
    ov.append([])
    ov.append(['Giả định'])
    for a in manifest['summary']['assumptions']:
        ov.append(['', a])
    ov.append([])
    ov.append(['File nguồn', 'sha256'])
    for f in manifest['summary']['sources']:
        ov.append([f['file'], f['sha256']])
    ov.column_dimensions['A'].width, ov.column_dimensions['B'].width = 70, 120

    inv = [i for i in state['invoices'] if i['period'] == PERIOD]
    sheet(wb, 'Hóa đơn', ['Mã HĐ', 'Phòng', 'Tòa', 'Trạng thái', 'Phải thu', 'Đã thu', 'Còn lại', 'Nguồn'],
          [[i['code'], code(rooms, i['roomId']), code(buildings, i['buildingId']), i['lifecycle'], i['totalDue'], paid.get(i['id'], 0), max(0, i['totalDue'] - paid.get(i['id'], 0)), ref(i)] for i in sorted(inv, key=lambda x: x['code'])],
          money=('Phải thu', 'Đã thu', 'Còn lại'))
    sheet(wb, 'Thu tiền', ['Mã phiếu', 'Ngày', 'Loại', 'Lượt thuê', 'Số tiền', 'Số hóa đơn phân bổ', 'Nguồn'],
          [[p['code'], p['receivedAt'], p['type'], code(stays, p.get('stayId')), p['amount'], len(p.get('allocations') or []), ref(p)] for p in sorted(state['payments'], key=lambda x: x['code'])], money=('Số tiền',))
    sheet(wb, 'Sổ cọc', ['Lượt thuê', 'Loại', 'Số tiền', 'Ngày', 'Kỳ', 'Ghi chú', 'Nguồn'],
          [[code(stays, l['stayId']), l['kind'], l['amount'], l.get('date'), l.get('period'), l.get('note', ''), ref(l)] for l in sorted(state['depositLedger'], key=lambda x: (code(stays, x['stayId']), x.get('date') or '', x['id']))], money=('Số tiền',))
    ex = [e for e in state['expenses'] if e['period'] == PERIOD]
    sheet(wb, 'Chi phí T9', ['Mã', 'Ngày', 'Loại', 'Dòng BC', 'Tòa / quỹ', 'Số tiền', 'Nguồn ghi', 'Ghi chú', 'Căn cứ'],
          [[e['code'], e['date'], e['category'], e.get('reportLine') or '', code(buildings, e.get('buildingId')) or e.get('fundCode') or '', e['amount'], e.get('source') or '', e.get('note') or '', ref(e)] for e in sorted(ex, key=lambda x: x['code'])], money=('Số tiền',))
    sheet(wb, 'Sửa chữa T9', ['Mã', 'Ngày', 'Tòa', 'Phòng', 'Nội dung', 'Tiền công', 'Vật tư', 'Trạng thái', 'Căn cứ'],
          [[r['code'], r['date'], r.get('buildingCode'), r.get('roomCode'), r.get('desc'), r.get('labor'), r.get('material'), r.get('status'), ref(r)] for r in state['repairLogs'] if r.get('period') == PERIOD], money=('Tiền công', 'Vật tư'))
    cm = [c for c in state['commissions'] if c.get('excel')]
    sheet(wb, 'Hoa hồng T9', ['Dòng Excel', 'Giao dịch', 'Phòng', 'Người nhận', 'Loại', 'F', 'H Excel', 'I duyệt', 'Trong vùng I3', 'Trạng thái', 'Đã chi', 'Cách ghi chi', 'Web còn thiếu'],
          [[c['excel']['row'], code(deals, c['dealId']), code(rooms, c['roomId']), c['recipient']['name'], c['recipient']['kind'], c['F'], c['excel']['H'], c.get('approvedAmount'), 'có' if c['excel']['inTotal'] else 'không',
            c['status'], sum(i['amount'] for i in c.get('installments') or []), 'theo Excel' if c.get('paidPerWorkbook') else ('trên web' if c.get('installments') else ''), ', '.join(c.get('eligibilityMissing') or [])]
           for c in sorted(cm, key=lambda x: x['excel']['row'])], money=('F', 'I duyệt', 'Đã chi'))
    run = next(r for r in state['payrollRuns'] if r['period'] == PERIOD and r['status'] == 'closed')
    sheet(wb, 'Lương T9', ['Mã NV', 'Nhân viên', 'Chức danh', 'Lương cơ bản', 'Ngày công', 'Thực lĩnh X'],
          [[code(emps, l['employeeId']), code(emps, l['employeeId'], 'name'), l.get('title'), l.get('base'), l.get('workdays'), l.get('X')] for l in run['lines']], money=('Lương cơ bản', 'Thực lĩnh X'))
    al = next(r for r in state['allocationRuns'] if r['period'] == PERIOD)
    sheet(wb, 'Phân bổ T9', ['Dòng', 'Quỹ', 'Tổng quỹ', 'Số tòa nhận', 'Tổng đã chia'],
          [[l['label'], l.get('fundCode'), l.get('fund'), len(l.get('results') or {}), sum((l.get('results') or {}).values())] for l in al['lines']], money=('Tổng quỹ', 'Tổng đã chia'))
    sr = next(r for r in state['shareRuns'] if r['period'] == PERIOD and r['buildingId'] == 'b_G1')
    sheet(wb, 'Bảng kê G1', ['Cổ đông', '% cổ phần', 'H', 'I', 'J', 'M (được nhận)'],
          [[code(shs, r['id']) or r['id'], r.get('pct'), r.get('H'), r.get('I'), r.get('J'), r.get('M')] for r in sr['rows']], money=('H', 'I', 'J', 'M (được nhận)'))
    msr = sorted([r for r in state.get('collectionMilestones', []) if r['period'] == PERIOD], key=lambda r: (code(buildings, r['buildingId']), r['milestone'], r.get('version') or 0))
    sheet(wb, 'Mốc thu T9', ['Tòa', 'Mốc (ngày)', 'Đã thu cộng dồn', 'Trạng thái', 'Người nhập', 'Người duyệt', 'Lý do từ chối', 'Nguồn'],
          [[code(buildings, r['buildingId']), r['day'], r['amount'], r['status'], r.get('reportedByName'), r.get('approvedBy') or r.get('rejectedBy') or '', r.get('rejectReason') or '', r.get('evidence') or ''] for r in msr],
          money=('Đã thu cộng dồn',))
    sheet(wb, 'Vấn đề', ['Loại', 'Tham chiếu', 'Số tiền', 'Trong vùng I3', 'Lý do'],
          [[i['kind'], i.get('ref') or i.get('id') or '', i.get('amount'), {True: 'có', False: 'không'}.get(i.get('inTotal'), ''), i.get('reason')] for i in manifest['summary']['issues']], money=('Số tiền',))
    wb.properties.creator = 'TimoHouse demo'
    wb.properties.created = wb.properties.modified = FIXED
    buf = io.BytesIO()
    wb.save(buf)
    # Ghi lại zip với thời gian cố định → cùng dữ liệu cho cùng byte.
    src = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            fixed = zipfile.ZipInfo(info.filename, date_time=(1980, 1, 1, 0, 0, 0))
            fixed.compress_type, fixed.external_attr = zipfile.ZIP_DEFLATED, info.external_attr
            data = src.read(info.filename)
            if info.filename == 'docProps/core.xml':  # openpyxl ghi giờ lưu vào "modified"
                data = re.sub(rb'(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)', rb'\g<1>2026-10-05T00:00:00Z\g<2>', data)
            dst.writestr(fixed, data)


def main():
    state = json.loads((ROOT / 'tmp' / 'september-flow' / 'state.json').read_text(encoding='utf-8'))
    manifest = json.loads((ROOT / 'docs' / 'data' / 'September_2026_Flow_Manifest.json').read_text(encoding='utf-8'))
    OUT.mkdir(parents=True, exist_ok=True)
    pdf, xlsx = OUT / 'contract-flow.pdf', OUT / 'Bo_du_lieu_mau_2026-09.xlsx'
    contract_pdf(state, manifest['summary']['journey'], pdf)
    workbook(state, manifest, xlsx)
    for p in (pdf, xlsx):
        print(f'  → {p.relative_to(ROOT)} ({os.path.getsize(p) / 1024:.0f} KB)')


if __name__ == '__main__':
    main()
