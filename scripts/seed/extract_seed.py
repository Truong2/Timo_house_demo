"""Sinh bộ dữ liệu demo Phase 1 từ file Excel của khách (chạy một lần, kết quả được commit).

    python -X utf8 scripts/seed/extract_seed.py

Nguồn (docs_timonouse/):
  SRC-08 Hoá đơn tiền nhà tháng 9 ... .xlsx      → tòa, phòng, lượt thuê, hóa đơn kỳ 2026-09, phiếu thu, phá HĐ, hoàn cọc
  SRC-04 bao_cao/BÁO CÁO KINH DOANH THÁNG 8.xlsx → tiền thuê nhà theo tòa, mức L1/L2/L3, benchmark báo cáo 08/2026
  SRC-03 bao_cao/bảng lương tháng 8.xlsx         → nhân viên, phân công, benchmark lương 08/2026
  SRC-06 Danh sách mã HĐ điện nước mạng.xlsx       → mã khách hàng nhà cung cấp theo tòa, hóa đơn nhà cung cấp 2026

Ẩn danh: mọi tên người được thay bằng bút danh xác định (cùng tên gốc → cùng bút danh), SĐT/CCCD/số TK sinh giả.
Giữ nguyên: mã tòa/phòng, giá, chỉ số, số tiền — là số liệu dùng để nghiệm thu.
Đầu ra: mockup/js/data/seed-*.js và tests/fixtures/*.json (file sinh, không sửa tay).
"""
import datetime as dt
import hashlib
import json
import os
import re
import sys
import unicodedata
import warnings

import openpyxl

warnings.filterwarnings('ignore')
ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..'))
SRC = os.path.join(ROOT, 'docs_timonouse')
OUT_JS = os.path.join(ROOT, 'mockup', 'js', 'data')
OUT_FX = os.path.join(ROOT, 'tests', 'fixtures')
PERIOD = '2026-09'
TODAY = '2026-09-29'


def log(*a):
    print(*a, flush=True)


def find(folder, pred):
    for x in os.listdir(folder):
        if pred(x):
            return os.path.join(folder, x)
    raise FileNotFoundError(folder)


def load(path):
    t = dt.datetime.now()
    wv = openpyxl.load_workbook(path, data_only=True)
    wf = openpyxl.load_workbook(path)
    log(f'  đọc {os.path.basename(path)} ({(dt.datetime.now() - t).seconds}s)')
    return wv, wf


def sheet(wb, name):
    for w in wb.worksheets:
        if w.title.strip() == name.strip():
            return w
    raise KeyError(name)


def col(letters):
    n = 0
    for ch in letters:
        n = n * 26 + (ord(ch) - 64)
    return n


def num(v):
    if v is None or v == '' or isinstance(v, str) and not v.strip():
        return 0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(',', '')
    try:
        return float(s)
    except ValueError:
        return 0


def r2(x):
    return round(float(x), 2)


def iso(v):
    if isinstance(v, dt.datetime):
        return v.strftime('%Y-%m-%d')
    return None


def norm(s):
    s = str(s or '').replace('Đ', 'D').replace('đ', 'd')
    s = unicodedata.normalize('NFD', s).encode('ascii', 'ignore').decode()
    return re.sub(r'\s+', ' ', s).strip().upper()


def h(key, mod=None):
    v = int(hashlib.md5(key.encode('utf-8')).hexdigest()[:12], 16)
    return v % mod if mod else v


def ekey(name):
    """Mã nhân viên ẩn danh: NV + 8 số băm từ tên chuẩn hóa – không suy ngược ra tên thật (id = emp_NV########).
    Không gọi pseudo() để không đổi thứ tự cấp tên giả hiển thị."""
    n = norm(name)
    return f'NV{h("emp:" + n, 10 ** 8):08d}' if n else ''


# ---------- bút danh ----------
HO = ['Nguyễn', 'Trần', 'Lê', 'Phạm', 'Hoàng', 'Huỳnh', 'Phan', 'Vũ', 'Võ', 'Đặng', 'Bùi', 'Đỗ', 'Hồ', 'Ngô', 'Dương', 'Lý', 'Mai', 'Trịnh']
DEM = ['Văn', 'Thị', 'Minh', 'Thu', 'Hoàng', 'Ngọc', 'Quang', 'Thanh', 'Đức', 'Hải', 'Anh', 'Gia', 'Bảo', 'Khánh', 'Phương', 'Tuấn', 'Diệu', 'Hữu']
TEN = ['An', 'Bình', 'Châu', 'Dũng', 'Giang', 'Hà', 'Hạnh', 'Hiếu', 'Hoa', 'Huy', 'Khánh', 'Lan', 'Linh', 'Long', 'Mai', 'Nam', 'Ngân', 'Nhung',
       'Phong', 'Phúc', 'Quân', 'Sơn', 'Tâm', 'Thảo', 'Trang', 'Trung', 'Tuấn', 'Vân', 'Vy', 'Yến', 'Đạt', 'Hùng', 'Loan', 'Minh', 'Oanh', 'Tú']
_used = {}
_reserved = None


def reserved():
    """Tên thật có thể xuất hiện trong mọi file nguồn (chủ tài khoản điện nước, khách, nhân viên, cổ đông, người nhận hoa hồng…):
    mọi cụm 3 từ liên tiếp của ô chữ trong các file .xlsx ở docs_timonouse, đã chuẩn hóa. Bút danh không được trùng một cụm nào (D19)."""
    global _reserved
    if _reserved is None:
        _reserved = set()
        for dp, _, fs in os.walk(SRC):
            for fn in fs:
                if not fn.lower().endswith('.xlsx') or fn.startswith('~$'):
                    continue
                try:
                    wb = openpyxl.load_workbook(os.path.join(dp, fn), read_only=True, data_only=True)
                except Exception:
                    continue
                for ws in wb.worksheets:
                    for row in ws.iter_rows(values_only=True):
                        for v in row:
                            if isinstance(v, str) and len(v) < 300:
                                w = [t for t in re.split(r'[^A-Z]+', norm(v)) if t]
                                for j in range(len(w) - 2):
                                    _reserved.add(' '.join(w[j:j + 3]))
                wb.close()
        log(f'  bút danh: {len(_reserved):,} cụm 3 từ trong file nguồn được loại trừ')
    return _reserved


def pseudo(key, scope='person'):
    k = scope + ':' + norm(key)
    if k in _used:
        return _used[k]
    i = 0
    while True:
        hv = h(k + '#' + str(i))
        name = f'{HO[hv % len(HO)]} {DEM[(hv // 7) % len(DEM)]} {TEN[(hv // 131) % len(TEN)]}'
        if name not in _used.values() and norm(name) not in reserved():  # D19: không trùng tên thật trong file nguồn
            _used[k] = name
            return name
        i += 1


def vendor_code(v):
    """Mã KH nhà cung cấp: che giữa; bỏ ô ghi chú kiểu 'hóa đơn', 'không có'."""
    if v is None:
        return None
    if isinstance(v, float):
        v = str(int(v))
    first = str(v).replace(',', ' ').split()[0] if str(v).strip() else ''
    if not re.fullmatch(r'[A-Za-z0-9]{6,}', first):
        return None
    return first[:4] + '•••' + first[-3:]


def phone(key):
    hv = h('phone:' + key)
    return '09' + str(hv % 10) + str(1000000 + hv % 8999999)[:7]


# ---------- SRC-08: hóa đơn tháng 9 ----------
NHA_COLS = dict(code='A', b='B', room='C', mgr='D', months='E', dh='F', gl='G', gm='H', gi='I', dep='J', kd='K', ld='L', od='M', ot='N', pp='O',
                el='P', wa='U', vs='Z', net='AC', elev='AF', ev='AI', wash='AL', combo='AO', common='AR', svc='AU', total='AV', note='AW', paid='AX',
                status='AY', moveIn='BA', endDate='BC')
PM_COLS = dict(code='A', b='B', room='C', mgr='D', months='E', gl='F', gm='G', gi='H', dep='I', kd='J', ld='K', od='L', ot='M', pp='N',
               el='O', wa='T', vs='Y', net='AB', elev='AE', ev='AH', wash='AK', combo='AN', svc='AQ', total='AR', note='AS', paid='AT', status='AU',
               moveIn='AW', endDate='AY')
LINE_KEYS = [('el', 3, True), ('wa', 4, True), ('vs', 5, False), ('net', 6, False), ('elev', 7, False), ('ev', 8, False), ('wash', 9, False),
             ('combo', 10, False)]
NHA_SHEETS = [('NHÀ T', 'T'), ('NHÀ S', 'S'), ('NHÀ G', 'G'), ('G16', 'G'), ('G17', 'G'), ('G18', 'G')]
DATE_RE = re.compile(r'(?<![\d/.])(\d{1,2})\s*[/.]\s*(\d{1,2})(?![\d/.])')


def parse_parts(formula, value):
    """AX '=1500000+2700000+4420000' → các lần thu; có số âm (chuyển phòng) → gộp một lần."""
    v = num(value)
    if isinstance(formula, str) and formula.startswith('='):
        body = formula[1:].replace(' ', '')
        if re.fullmatch(r'[\d.+\-]+', body):
            parts = [float(x) for x in re.findall(r'[+-]?[\d.]+', body)]
            if parts and all(p > 0 for p in parts) and abs(sum(parts) - v) < 1:
                return parts
    return [v] if v else []


def note_dates(note, month):
    """Ngày trong ghi chú: 'kh mới 6/9' = ngày vào, '5.9 tt đủ' = ngày thu."""
    s = str(note or '')
    move, pay = None, []
    for m in DATE_RE.finditer(s):
        d, mo = int(m.group(1)), int(m.group(2))
        if not (1 <= d <= 31 and 1 <= mo <= 12):
            continue
        before = s[max(0, m.start() - 8):m.start()].lower()
        if 'mới' in before or 'từ' in before or 'vào' in before:
            move = move or (mo, d)
        elif mo in (month, month - 1):
            pay.append((mo, d))
    return move, pay


def read_row(sv, sf, r, C):
    g = lambda k: sv.cell(r, col(C[k])).value if k in C else None
    gf = lambda k: sf.cell(r, col(C[k])).value if k in C else None
    trip = lambda k: [num(sv.cell(r, col(C[k]) + i).value) for i in range(3)]
    five = lambda k: [num(sv.cell(r, col(C[k]) + i).value) for i in range(5)]
    return dict(code=str(g('code') or '').strip(), b=str(g('b') or '').strip().upper(), room=g('room'), mgr=str(g('mgr') or '').strip(),
                months=num(g('months')), dh=num(g('dh')), gl=num(g('gl')), gm=num(g('gm')), gi=num(g('gi')), gi_f=gf('gi'), gm_f=gf('gm'),
                dep=num(g('dep')), kd=num(g('kd')), ld=num(g('ld')), od=num(g('od')), ot=num(g('ot')), ot_f=gf('ot'), pp=num(g('pp')),
                el=five('el'), wa=five('wa'), vs=trip('vs'), net=trip('net'), elev=trip('elev'), ev=trip('ev'), wash=trip('wash'),
                combo=trip('combo'), common=trip('common') if 'common' in C else [0, 0, 0], svc=num(g('svc')), svc_f=gf('svc'),
                total=num(g('total')), total_f=gf('total'), note=str(g('note') or '').strip(), paid=num(g('paid')), paid_f=gf('paid'),
                status=str(g('status') or '').strip(), moveIn=iso(g('moveIn')), endDate=iso(g('endDate')))


def svc_factor(formula):
    m = re.search(r'/\s*(\d+)\s*\)\s*\*\s*(\d+)', str(formula or ''))
    return (int(m.group(2)), int(m.group(1))) if m else None


def rent_factor(formula):
    m = re.search(r'\(\s*([\d.]+)\s*/\s*(\d+)\s*\)\s*\*\s*(\d+)', str(formula or ''))
    return (int(m.group(3)), int(m.group(2)), float(m.group(1))) if m else None


def build_lines(x, kind):
    """13 dòng theo mẫu in: [no, chỉ số mới, chỉ số cũ, SL, hệ số, đơn giá, thành tiền]."""
    L = []
    if kind == 'PM':  # PHÒNG MỚI: H = (giá/30)×ngày — tiền tháng lẻ đã quy đổi sẵn
        rf = rent_factor(x['gi_f'])
        fac = round(x['gi'] / x['monthly'], 6) if x['monthly'] else 1
        L.append([1, None, None, x['months'] or 1, fac, x['monthly'], r2(x['gi'] * (x['months'] or 1))])
        x['rentDays'] = [rf[0], rf[1]] if rf else None
    else:
        L.append([1, None, None, x['months'], 1, x['gi'], r2(x['gi'] * x['months'])])
    L.append([2, None, None, 1 if x['dep'] else 0, 1, x['dep'], x['dep']])
    sf = svc_factor(x['svc_f']) if kind == 'PM' else None
    fac = (sf[0] / sf[1]) if sf else 1
    for key, no, meter in LINE_KEYS:
        v = x[key]
        if meter:
            prev, curr, qty, unit, amt = v
            if not prev and not curr:  # nước theo người (không đồng hồ) = dịch vụ cố định → chia ngày như DV
                L.append([no, None, None, qty, round(fac, 6), unit, r2(amt * fac)])
            else:  # theo chỉ số: không chia ngày
                L.append([no, curr or None, prev or None, qty, 1, unit, r2(amt)])
        else:
            qty, unit, amt = v
            L.append([no, None, None, qty, round(fac, 6), unit, r2(amt * fac)])
    L.append([11, None, None, 1 if x['od'] else 0, 1, x['od'], x['od']])
    q, u, a = x['common']
    L.append([12, None, None, q, 1, u, r2(a)])
    L.append([13, None, None, 1 if x['ot'] else 0, 1, r2(x['ot']), r2(x['ot'])])
    return L, sf


def main():
    os.makedirs(OUT_JS, exist_ok=True)
    os.makedirs(OUT_FX, exist_ok=True)
    f08 = find(SRC, lambda x: 'tháng 9' in x and 'dịch' in x and x.endswith('.xlsx'))
    f04 = find(os.path.join(SRC, 'bao_cao'), lambda x: 'KINH DOANH' in x.upper() and x.endswith('.xlsx'))
    f03 = find(os.path.join(SRC, 'bao_cao'), lambda x: 'lương' in x and x.endswith('.xlsx'))
    f06 = find(SRC, lambda x: x.startswith('Danh s') and x.endswith('.xlsx'))
    wv8, wf8 = load(f08)

    rows = []
    for name, grp in NHA_SHEETS:
        sv, sf = sheet(wv8, name), sheet(wf8, name)
        for r in range(5, sv.max_row + 1):
            x = read_row(sv, sf, r, NHA_COLS)
            if not x['code'] or not x['b'] or x['code'].upper() in ('MÃ',):
                continue
            x.update(sheet=name, grp=grp, row=r, kind='NHA', monthly=x['gi'], list=x['gl'])
            rows.append(x)
    for name, kind in [('PHÒNG MỚI THÁNG 9', 'PM9'), ('PHÒNG MỚI THÁNG 10', 'PM10')]:
        sv, sf = sheet(wv8, name), sheet(wf8, name)
        for r in range(5, sv.max_row + 1):
            x = read_row(sv, sf, r, PM_COLS)
            if not x['code'] or not x['b']:
                continue
            # PHÒNG MỚI: F niêm yết, G giá thuê tháng, H tiền phòng tháng lẻ đã chia ngày
            x.update(sheet=name, grp=None, row=r, kind=kind, monthly=x['gm'] or x['gl'], list=x['gl'] or x['gm'])
            rows.append(x)
    log(f'  dòng hóa đơn đọc được: {len(rows)}')

    # --- phá HĐ ---
    sv = sheet(wv8, 'DS phòng phá hđ')
    breach = {}
    for r in range(4, sv.max_row + 1):
        code = str(sv.cell(r, 2).value or '').strip().upper()
        if not code:
            continue
        breach[code] = dict(code=code, mgr=str(sv.cell(r, 3).value or '').strip(), start=iso(sv.cell(r, 4).value),
                            months=num(sv.cell(r, 5).value) or None, reason=str(sv.cell(r, 6).value or '').strip(),
                            due=num(sv.cell(r, 8).value), paid=num(sv.cell(r, 9).value))
    log(f'  phá HĐ: {len(breach)}')

    # --- điện nước phòng trống ---
    sv = sheet(wv8, 'ĐIỆN NƯỚC PHÒNG TRỐNG')
    vacantUtil = []
    for r in range(5, sv.max_row + 1):
        code = str(sv.cell(r, 1).value or '').strip()
        if not code:
            continue
        v = [num(sv.cell(r, c).value) for c in range(2, 13)]
        vacantUtil.append(dict(room=code, elPrev=v[0], elCurr=v[1], elQty=v[2], elPrice=v[3], elAmt=v[4], note=str(sv.cell(r, 7).value or '').strip(),
                               waPrev=v[6], waCurr=v[7], waQty=v[8], waPrice=v[9], waAmt=v[10]))

    # --- hoàn cọc ---
    sv, sf = sheet(wv8, 'HOÀN CỌC'), sheet(wf8, 'HOÀN CỌC')
    refundsRaw = []
    RC = [('S', 'electric', 'O'), ('X', 'water', 'T'), ('AA', 'cleaningSvc', None), ('AD', 'internet', None), ('AG', 'elevator', None), ('AJ', 'ev', None),
          ('AM', 'washer', None), ('AP', 'depreciation', None), ('AS', 'repair', None), ('AV', 'cleaning', None), ('AY', 'painting', None), ('AZ', 'other', None)]
    for r in range(5, sv.max_row + 1):
        code = str(sv.cell(r, 1).value or '').strip()
        if not code or not sv.cell(r, 2).value:
            continue
        dedu = []
        for amtCol, kind, meterCol in RC:
            c = col(amtCol)
            if kind == 'other':
                qty, unit, amt = num(sv.cell(r, c).value), num(sv.cell(r, c + 1).value), num(sv.cell(r, c + 2).value)
            else:
                qty, unit, amt = num(sv.cell(r, c - 2).value), num(sv.cell(r, c - 1).value), num(sv.cell(r, c).value)
            if meterCol:
                mc = col(meterCol)
                prev, curr = num(sv.cell(r, mc).value), num(sv.cell(r, mc + 1).value)
                qty = num(sv.cell(r, mc + 2).value)
                unit = num(sv.cell(r, mc + 3).value)
                amt = num(sv.cell(r, mc + 4).value)
            else:
                prev = curr = None
            if amt:
                dedu.append(dict(kind=kind, prev=prev, curr=curr, qty=qty, unit=unit, amount=r2(amt)))
        extra = rent_factor(sf.cell(r, 8).value)
        refundsRaw.append(dict(room=code, b=str(sv.cell(r, 2).value).strip().upper(), mgr=str(sv.cell(r, 4).value or '').strip(), deposit=num(sv.cell(r, 9).value),
                               deductions=dedu, bc=num(sv.cell(r, col('BC')).value), bd=num(sv.cell(r, col('BD')).value),
                               status=str(sv.cell(r, col('BE')).value or '').strip(), note=str(sv.cell(r, col('BF')).value or '').strip(),
                               extraDays=extra[0] if extra else 0, extraDaysDenom=extra[1] if extra else None, extraRent=r2(num(sv.cell(r, 8).value)), row=r))
    log(f'  hoàn cọc: {len(refundsRaw)}')

    # --- SRC-04: báo cáo theo tòa 08/2026 ---
    wv4, wf4 = load(f04)
    report = dict(total={}, business={}, byBuilding={}, level={}, labels={})
    st, sk = sheet(wv4, 'BÁO CÁO TỔNG THÁNG 8'), sheet(wv4, 'BÁO CÁO KINH DOANH THÁNG 8')
    for r in range(3, 62):
        report['labels'][r] = [str(st.cell(r, 1).value or '').strip(), str(st.cell(r, 2).value or '').strip()]
        report['total'][r] = [num(st.cell(r, c).value) for c in (3, 4, 5, 6)]
        report['business'][r] = [num(sk.cell(r, c).value) for c in (3, 4, 5, 6)]
    for name, grp in [('BC DT THÁNG 8 NHÀ T', 'T'), ('BCDT THÁNG 8 NHÀ S', 'S'), ('BC DT THÁNG 8 NHÀ G', 'G')]:
        s = sheet(wv4, name)
        for c in range(4, s.max_column + 1):
            code = str(s.cell(2, c).value or '').strip().upper()  # mã tòa in hoa (nguồn ghi lẫn s8/S8)
            if not code:
                continue
            report['level'][code] = str(s.cell(3, c).value or '').strip() or None
            vals = {}
            for r in range(4, 49):
                v = s.cell(r, c).value
                if isinstance(v, (int, float)) and v:
                    vals[r - 1] = r2(v) if abs(v) > 1 else v  # excelRow theo sheet tổng
            report['byBuilding'][code] = dict(grp=grp, v=vals)
    log(f'  benchmark báo cáo: {len(report["byBuilding"])} tòa')

    # --- SRC-06: mã nhà cung cấp + hóa đơn 2026 ---
    wv6 = openpyxl.load_workbook(f06, data_only=True)
    s6 = sheet(wv6, 'năm 2026')
    vendors = {}
    for r in range(4, s6.max_row + 1):
        code = str(s6.cell(r, 1).value or '').strip()
        if not code:
            continue
        mask = vendor_code
        bills = lambda c0, n: {f'2026-{i + 1:02d}': num(s6.cell(r, c0 + i).value) for i in range(n) if num(s6.cell(r, c0 + i).value)}
        vendors[code] = dict(
            electric=dict(code=mask(s6.cell(r, 2).value), holder=pseudo(s6.cell(r, 6).value or code, 'holder') if s6.cell(r, 6).value else None, bills=bills(col('G'), 8)),
            water=dict(code=mask(s6.cell(r, col('O')).value), holder=pseudo(s6.cell(r, col('P')).value or code, 'holder') if s6.cell(r, col('P')).value else None, bills=bills(col('Q'), 8)),
            internet=dict(code=mask(s6.cell(r, col('Y')).value), holder=pseudo(s6.cell(r, col('Z')).value or code, 'holder') if s6.cell(r, col('Z')).value else None, bills=bills(col('AA'), 3)),
            note=str(s6.cell(r, col('AH')).value or '').strip())
    log(f'  nhà cung cấp: {len(vendors)} tòa')

    # --- SRC-03: bảng lương 08/2026 ---
    wv3, wf3 = load(f03)
    s3v, s3f = sheet(wv3, 'THÁNG 8'), sheet(wf3, 'THÁNG 8')
    payroll, cur = [], None
    for r in range(10, 164):
        name = s3v.cell(r, 2).value
        title = s3v.cell(r, 3).value
        bcode = s3v.cell(r, 9).value
        if name and title:
            cur = dict(key=ekey(name), name=pseudo(name), title=str(title).strip(), row=r,
                       base=num(s3v.cell(r, 4).value), lunch=num(s3v.cell(r, 5).value), fuel=num(s3v.cell(r, 6).value),
                       lead=num(s3v.cell(r, 7).value), support=num(s3v.cell(r, 8).value), total=num(s3v.cell(r, 23).value),
                       net=num(s3v.cell(r, 24).value), leadF=str(s3f.cell(r, 7).value or ''), baseF=str(s3f.cell(r, 4).value or ''), buildings=[])
            payroll.append(cur)
        if cur and isinstance(bcode, str) and bcode.strip():
            vals = {k: num(s3v.cell(r, col(k)).value) for k in 'JKLMNOPQRSTUVW'}
            vals['Lf'] = str(s3f.cell(r, col('L')).value or '')
            vals['Mf'] = str(s3f.cell(r, col('M')).value or '')
            vals['Vf'] = str(s3f.cell(r, col('V')).value or '')
            vals['Sf'] = str(s3f.cell(r, col('S')).value or '')
            vals['row'] = r
            cur['buildings'].append(dict(b=bcode.strip().upper(), **vals))
    log(f'  nhân viên bảng lương: {len(payroll)}; dòng tòa: {sum(len(p["buildings"]) for p in payroll)}')

    # ---------- chuẩn hóa thực thể ----------
    grpOf = {}
    for x in rows:
        if x['grp']:
            grpOf.setdefault(x['b'], x['grp'])
    for code, v in report['byBuilding'].items():
        grpOf.setdefault(code, v['grp'])
    AREAS = [('A1', 'Khu Cầu Giấy'), ('A2', 'Khu Nam Từ Liêm'), ('A3', 'Khu Thanh Xuân'), ('A4', 'Khu Hà Đông')]
    STREETS = ['Trần Thái Tông', 'Xuân Thủy', 'Hồ Tùng Mậu', 'Mỹ Đình', 'Nguyễn Trãi', 'Khuất Duy Tiến', 'Quang Trung', 'Phùng Khoang', 'Cầu Giấy', 'Lê Đức Thọ']

    # nhân viên
    emps = {}
    for p in payroll:
        emps[p['key']] = dict(key=p['key'], name=p['name'], title=p['title'].upper().replace('KĨ', 'KỸ'))
    for x in rows:
        k = ekey(x['mgr'])
        if k and k not in emps:
            emps[k] = dict(key=k, name=pseudo(x['mgr']), title='NVVH')

    # tòa
    bcodes = sorted({x['b'] for x in rows} | set(report['byBuilding']), key=lambda c: (c[0], int(re.sub(r'\D', '', c) or 0), c))
    mgrOfB, mgrReal = {}, {}
    for x in rows:
        if x['kind'] == 'NHA' and x['mgr']:
            mgrOfB.setdefault(x['b'], ekey(x['mgr']))
            mgrReal.setdefault(x['b'], norm(x['mgr']))  # chỉ dùng nội bộ để chọn mẫu hóa đơn, không ghi ra file
    buildings = []
    for c in bcodes:
        grp = grpOf.get(c, c[0] if c[0] in 'TSG' else 'S')
        bb = report['byBuilding'].get(c, {}).get('v', {})
        a = AREAS[h('area:' + c, len(AREAS))]
        buildings.append(dict(code=c, group=grp, areaId=a[0], level=report['level'].get(c), ownerRent=bb.get(20, 0) or 0,
                              address=f'Số {10 + h("addr:" + c, 180)} ngõ {h("ngo:" + c, 300) + 1} {STREETS[h("st:" + c, len(STREETS))]}, Hà Nội',
                              managerKey=mgrOfB.get(c), vendor=vendors.get(c),
                              template={'T': 'TECH', 'G': 'G1_TECH'}.get(grp, 'VP_HANG' if mgrReal.get(c, '').startswith('NGUYEN THI THUONG') else 'VP')))

    # phòng + lượt thuê + hóa đơn
    rooms, stays, invoices, payments, depositLedger, customers = {}, [], [], [], [], []
    seq = {}

    def next_code(room):
        seq[room] = seq.get(room, 0) + 1
        return f'{room}A{seq[room]:03d}'

    def add_room(code, b, number, x=None):
        if code in rooms:
            return rooms[code]
        # số phòng: cột số (có thể là chữ như "501T2") → số đầu của mã phòng; không có số → 0
        try:
            n = int(float(number))
        except (TypeError, ValueError):
            n = int(re.sub(r'\D.*', '', str(code)) or 0)
        price = x['monthly'] if x else 0
        # đồng hồ chung / dòng thu chung của tòa: mã 000…/001… hoặc không có số phòng và không có giá thuê
        common = bool(re.match(r'^00\d', code)) or (n == 0 and not price)
        rm = dict(code=code, b=b, number=n, floor=n // 100 if n >= 100 else 0, list=x['list'] if x else 0, mgmt=x['gm'] if x else 0,
                  price=price, exploitation='meter_common' if common else 'timehouse', status='vacant_ready')
        rooms[code] = rm
        return rm

    def mk_customer(stayCode, pp):
        cid = 'KH-' + stayCode
        occ = 'Sinh viên' if h('occ:' + stayCode, 10) < 4 else 'Người đi làm'
        customers.append(dict(id=cid, name=pseudo(stayCode, 'customer'), phone=phone(stayCode), occupation=occ,
                              idNo='0' + str(10000000000 + h('cccd:' + stayCode, 89999999999))[:11], zalo=h('zl:' + stayCode, 10) < 8))
        return cid

    def dates_for(x, code):
        start = x['moveIn']
        if not start:
            months_ago = 1 + h('start:' + code, 26)
            y, m = 2026, 8 - months_ago
            while m <= 0:
                m += 12; y -= 1
            start = f'{y}-{m:02d}-{1 + h("d:" + code, 27):02d}'
        end = x['endDate']
        if not end:
            y, m, d = map(int, start.split('-'))
            y2, m2 = (y + 1, m) if h('term:' + code, 3) else (y, m + 6) if m <= 6 else (y + 1, m - 6)
            end = (dt.date(y2, m2, min(d, 28)) - dt.timedelta(days=1)).isoformat()
        while end < TODAY and x['kind'] == 'NHA':
            y, m, d = map(int, end.split('-'))
            end = dt.date(y + 1, m, min(d, 28)).isoformat()
        return start, end

    # ưu tiên: lượt thuê đã kết thúc (hoàn cọc) → dòng NHÀ → phòng mới T9 → phòng mới T10
    refundStayByRoom = {}
    for rr in refundsRaw:
        rm = add_room(rr['room'], rr['b'], None)
        code = next_code(rr['room'])
        end = '2026-08-31'
        if rr['extraDays']:
            end = f'2026-09-{rr["extraDays"]:02d}'
        cid = mk_customer(code, 1)
        stays.append(dict(code=code, room=rr['room'], b=rr['b'], customerId=cid, status='ended', endType='expired', mgr=ekey(rr['mgr']),
                          moveIn=None, rentStart=None, svcStart=None, endDate=end, noticeDate=None, deposit=rr['deposit'],
                          depositStatus='refunded' if 'hoàn' in rr['status'].lower() else 'refund_pending', rent=0, people=1, source='HOÀN CỌC'))
        refundStayByRoom.setdefault(rr['room'], []).append(code)
        rr['stay'] = code

    breachUsed = set()
    issueDate, dueFrom, dueTo = '2026-08-22', '2026-08-25', '2026-08-31'
    for x in rows:
        code_room = x['code']
        rm = add_room(code_room, x['b'], x['room'], x)
        if x['kind'] == 'NHA' and x['total'] == 0 and x['paid'] == 0:
            rm['status'] = 'vacant_ready'
            rm['srcStatus'] = x['status'] or 'Trống'
            continue
        stayCode = next_code(code_room)
        cid = mk_customer(stayCode, x['pp'])
        start, end = dates_for(x, stayCode)
        status, endType, depStatus = 'active', None, 'held'
        move, payDates = note_dates(x['note'], 9 if x['kind'] != 'PM10' else 10)
        isBreach = x['kind'] == 'NHA' and code_room in breach and code_room not in breachUsed and (x['paid'] < x['total'] or x['status'] in ('Chưa TT', 'Thiếu'))
        if isBreach:
            breachUsed.add(code_room)
            bi = breach[code_room]
            status, depStatus = 'ended', 'kept_breach'
            endType = 'abscond' if 'trốn' in bi['reason'].lower() else 'breach'
            if bi['start']:
                start = bi['start']
            end = '2026-09-05'
        if x['kind'] == 'PM9':
            if move:
                start = f'2026-{move[0]:02d}-{move[1]:02d}'
            else:
                rf = rent_factor(x['gi_f'])
                start = f'2026-09-{31 - rf[0]:02d}' if rf else '2026-09-01'
            y, m, d = map(int, start.split('-'))
            end = (dt.date(y + 1, m, d) - dt.timedelta(days=1)).isoformat()
        if x['kind'] == 'PM10':
            status = 'pending'
            start = f'2026-10-{move[1]:02d}' if move and move[0] == 10 else '2026-10-01'
            end = (dt.date(2027, 10, int(start[-2:])) - dt.timedelta(days=1)).isoformat()
        dep = x['dh'] or x['dep']
        if x['kind'] in ('PM9', 'PM10'):
            dep = x['dep']
        mgr = ekey(x['mgr'])
        stays.append(dict(code=stayCode, room=code_room, b=x['b'], customerId=cid, status=status, endType=endType, mgr=mgr,
                          moveIn=start, rentStart=start, svcStart=start, endDate=end, deposit=dep, depositStatus=depStatus,
                          rent=x['monthly'], list=x['list'], people=x['pp'] or 1,
                          breachReason=breach[code_room]['reason'] if isBreach else None, source=x['sheet'], row=x['row'],
                          payMonths=x['months'] or 1, newDeposit=x['dep']))
        if status in ('active',):
            rm['status'] = 'occupied'
        if status == 'pending' and rm['status'] != 'occupied':
            rm['status'] = 'reserved'
        if rm['price'] == 0 and x['monthly']:
            rm.update(list=x['list'], mgmt=x['gm'], price=x['monthly'])

        # --- hóa đơn ---
        if x['kind'] == 'PM10':
            # khách đã cọc/đóng trước cho tháng 10: tiền nhận ghi phiếu thu trả trước + cọc
            parts = parse_parts(x['paid_f'], x['paid'])
            depAmt = min(x['dep'], sum(parts))
            for i, p in enumerate(parts):
                payments.append(dict(stay=stayCode, invoice=None, kind='deposit' if i == 0 and depAmt else 'prepay', amount=p, date=None,
                                     dateSource='synthetic', period='2026-10', part=i + 1))
            if depAmt:
                depositLedger.append(dict(stay=stayCode, kind='receive', amount=depAmt, period='2026-09'))
            continue
        kind = 'PM' if x['kind'] == 'PM9' else 'NHA'
        lines, sf = build_lines(x, kind)
        inv = dict(code=f'INV-2026-09-{stayCode}', stay=stayCode, room=code_room, b=x['b'], period='2026-09', issueDate=issueDate,
                   dueFrom=dueFrom, dueTo=dueTo, cutoff='2026-08-22', lifecycle='issued', isNewStay=x['kind'] == 'PM9', isBreach=isBreach,
                   lines=[l for l in lines if l[0] in (1, 3, 4) or l[3] or l[6]], excelTotal=r2(x['total']), excelPaid=r2(x['paid']), excelStatus=x['status'], note=x['note'],
                   hardTotal=not (isinstance(x['total_f'], str) and x['total_f'].startswith('=')),
                   svcFactor=list(sf) if sf else None, src=f"{x['sheet']}!{x['row']}", months=x['months'] or 1)
        if isBreach:
            inv['breachDue'] = breach[code_room]['due']
            inv['breachPaid'] = breach[code_room]['paid']
        invoices.append(inv)
        parts = parse_parts(x['paid_f'], x['paid'])
        for i, p in enumerate(parts):
            d = None
            if payDates:
                mo, dd = payDates[min(i, len(payDates) - 1)]
                d = f'2026-{mo:02d}-{dd:02d}'
            payments.append(dict(stay=stayCode, invoice=inv['code'], kind='invoice', amount=p, date=d, dateSource='note' if d else 'synthetic',
                                 period='2026-09', part=i + 1))
        if x['dep']:
            depositLedger.append(dict(stay=stayCode, kind='receive', amount=x['dep'], period='2026-09'))
        if isBreach and dep:
            depositLedger.append(dict(stay=stayCode, kind='keep_breach', amount=dep, period='2026-09'))

    # --- ngày thu tổng hợp: khớp lũy kế 89,0% / 97,9% / 98,4% tại mốc 5/10/15 ---
    invPay = [p for p in payments if p['invoice']]
    total = sum(p['amount'] for p in invPay)
    target = [('2026-09-05', 0.890), ('2026-09-10', 0.979), ('2026-09-15', 0.984), ('2026-09-28', 1.0)]
    known = sorted([p for p in invPay if p['date']], key=lambda p: p['date'])
    unknown = sorted([p for p in invPay if not p['date']], key=lambda p: h('pd:' + p['stay'] + str(p['part'])))
    cum = {t[0]: sum(p['amount'] for p in known if p['date'] <= t[0]) for t in target}
    pool = list(unknown)
    prev_day = '2026-08-24'
    for day, frac in target:
        need = frac * total - cum[day]
        span_start = dt.date.fromisoformat(prev_day) + dt.timedelta(days=1)
        span = (dt.date.fromisoformat(day) - span_start).days + 1
        while pool and need > 0:
            p = pool.pop(0)
            off = span - 1 if h('w:' + p['stay'], 100) < 45 else h('o:' + p['stay'], span)
            p['date'] = (span_start + dt.timedelta(days=off)).isoformat()
            need -= p['amount']
            for d2, _ in target:
                if p['date'] <= d2:
                    cum[d2] += p['amount']
        prev_day = day
    for p in pool:
        p['date'] = '2026-09-20'
    for p in payments:
        if not p['date']:
            p['date'] = '2026-09-' + f'{10 + h("pp:" + p["stay"], 18):02d}'
    for i, p in enumerate(sorted(payments, key=lambda p: (p['date'], p['stay'], p['part']))):
        p['code'] = f'PT-{p["date"][:7].replace("-", "")}-{i + 1:04d}'
        hm = h('m:' + p['stay'] + str(p['part']), 100)
        p['method'] = 'cash' if hm < 12 else 'bank'

    # --- phiếu hoàn cọc ---
    refunds = []
    for i, rr in enumerate(refundsRaw):
        done = 'hoàn' in rr['status'].lower()
        refunds.append(dict(code=f'HC-202609-{i + 1:04d}', stay=rr['stay'], room=rr['room'], b=rr['b'], deposit=rr['deposit'], deductions=rr['deductions'],
                            excelBC=rr['bc'], excelBD=rr['bd'], extraDays=rr['extraDays'], extraDaysDenom=rr['extraDaysDenom'], extraRent=rr['extraRent'],
                            status='paid' if done else 'calculated', note=rr['note'], src=f'HOÀN CỌC!{rr["row"]}'))
        depositLedger.append(dict(stay=rr['stay'], kind='refund' if done else 'deduct', amount=rr['bd'] if done else 0, period='2026-09'))

    # --- phân công (từ bảng lương T8, ghi đè bằng quản lý tháng 9 nếu khác) ---
    assignments = []
    augOwner = {}
    for p in payroll:
        for b in p['buildings']:
            augOwner.setdefault(b['b'], p['key'])
    for bld in buildings:
        c = bld['code']
        aug, sep = augOwner.get(c), bld['managerKey']
        if aug and sep and aug != sep:
            assignments.append(dict(emp=aug, b=c, resp='operate', start='2026-01-01', end='2026-08-31'))
            assignments.append(dict(emp=sep, b=c, resp='operate', start='2026-09-01', end=None))
        elif aug or sep:
            assignments.append(dict(emp=aug or sep, b=c, resp='operate', start='2026-01-01', end=None))

    # ---------- xuất ----------
    master = dict(generatedFrom=[os.path.basename(f) for f in (f08, f04, f03, f06)], today=TODAY, areas=[dict(id=a, name=n) for a, n in AREAS],
                  buildings=buildings, rooms=list(rooms.values()), employees=list(emps.values()), assignments=assignments,
                  customers=customers, stays=stays)
    p09 = dict(period=PERIOD, invoices=invoices, payments=payments, depositLedger=depositLedger, refunds=refunds, vacantUtil=vacantUtil,
               breach=[dict(code=b['code'], reason=b['reason'], due=b['due'], paid=b['paid'], start=b['start'], months=b['months']) for b in breach.values()])
    bench = dict(period='2026-08', report=report, payroll=payroll, vendors={k: v for k, v in vendors.items()})

    def write_js(name, var, obj):
        path = os.path.join(OUT_JS, name)
        txt = json.dumps(obj, ensure_ascii=False, separators=(',', ':'))
        with open(path, 'w', encoding='utf-8', newline='\n') as f:
            f.write(f'/* FILE SINH TỰ ĐỘNG bởi scripts/seed/extract_seed.py – không sửa tay. Dữ liệu đã ẩn danh. */\n')
            f.write(f'window.TH = window.TH || {{}}; TH.data = TH.data || {{}}; TH.data.{var} = {txt};\n')
        log(f'  → {os.path.relpath(path, ROOT)} ({len(txt) / 1024:.0f} KB)')

    write_js('seed-master.js', 'master', master)
    write_js('seed-2026-09.js', 'p202609', p09)
    write_js('seed-2026-08-bench.js', 'bench202608', bench)
    for name, obj in [('invoices-2026-09.json', invoices), ('report-2026-08.json', report), ('payroll-2026-08.json', payroll),
                      ('refunds-2026-09.json', refunds), ('breach-2026-09.json', p09['breach'])]:
        with open(os.path.join(OUT_FX, name), 'w', encoding='utf-8', newline='\n') as f:
            json.dump(obj, f, ensure_ascii=False, indent=0)
    # tóm tắt
    by = {}
    for inv in invoices:
        by[inv['excelStatus']] = by.get(inv['excelStatus'], 0) + 1
    log(f'  tòa {len(buildings)} · phòng {len(rooms)} · lượt thuê {len(stays)} · hóa đơn {len(invoices)} · phiếu thu {len(payments)} · hoàn cọc {len(refunds)}')
    log(f'  trạng thái Excel: {by}')
    pm9 = sum(p['amount'] for p in payments if p['invoice'] and any(i['code'] == p['invoice'] and i['isNewStay'] for i in invoices))
    log(f'  tổng đã thu PHÒNG MỚI T9 = {pm9:,.0f}')


if __name__ == '__main__':
    main()
