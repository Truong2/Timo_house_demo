"""Đọc bằng chứng Excel cho bộ dữ liệu luồng tháng 9 (không đổi seed lịch sử, không lộ tên người).

    python -X utf8 scripts/seed/extract_september_flow.py

Đầu ra: scripts/seed/september-source.json, gồm
  - files: danh sách file Excel nguồn + sha256 + tên sheet (truy vết);
  - commissions: sheet HOA HỒNG THÁNG 9.26 (SRC-09), cùng cấu trúc dòng với T8 (extract_seed_p2.commission) đã ẩn danh.
"""
from pathlib import Path
import hashlib
import json
import os
import sys

import openpyxl

sys.path.insert(0, os.path.dirname(__file__))
from extract_seed_p2 import commission, mask_people  # noqa: E402

ROOT = Path(__file__).resolve().parents[2]
SHEET = 'HOA HỒNG THÁNG 9.26'
TOTAL_FROM_ROW = 14  # I3 = SUBTOTAL(109, I14:I2003): Excel bỏ qua dòng 4…13


def main():
    files = []
    for path in sorted((ROOT / 'docs_timonouse').rglob('*.xlsx')):
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        files.append({'file': str(path.relative_to(ROOT)).replace('\\', '/'), 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(), 'sheets': wb.sheetnames})
        if path.name.startswith('Hoa hồng'):
            last_row = max(r for r, row in enumerate(wb[SHEET].iter_rows(min_col=4, max_col=9, values_only=True), 1) if any(v is not None for v in row))
        wb.close()
    cm = commission(SHEET, '2026-09', status_col=2, ql_col=3, last_row=last_row)
    cm = json.loads(mask_people(json.dumps(cm, ensure_ascii=False)))
    with_i = [r for r in cm['rows'] if r['I'] is not None]
    cm['totalFromRow'] = TOTAL_FROM_ROW
    cm['sumAll'] = round(sum(r['I'] for r in with_i), 2)
    cm['sumInTotalRange'] = round(sum(r['I'] for r in with_i if r['row'] >= TOTAL_FROM_ROW), 2)
    assert abs(cm['sumInTotalRange'] - cm['excelTotal']) < 1, (cm['sumInTotalRange'], cm['excelTotal'])
    print(f"  Σ mọi dòng = {cm['sumAll']:,.2f} · Σ từ dòng {TOTAL_FROM_ROW} = {cm['sumInTotalRange']:,.2f} = I3 {cm['excelTotal']:,.2f}")
    out = ROOT / 'scripts' / 'seed' / 'september-source.json'
    out.write_text(json.dumps({'period': '2026-09', 'files': files, 'commissions': cm}, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print('Wrote', out.relative_to(ROOT))


if __name__ == '__main__':
    main()
